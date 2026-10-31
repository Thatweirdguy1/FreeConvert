"""Reproduce the usability and data-preservation failures with synthetic data."""

import os
import sqlite3
import tempfile
import time
import threading
import unittest
from contextlib import closing
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PIL import Image
from PySide6.QtCore import QPointF, Qt, QCoreApplication, QEvent
from PySide6.QtGui import QMouseEvent
from PySide6.QtWidgets import QApplication, QFrame, QMessageBox, QPushButton, QTableWidgetItem
from app import Window, preferred_format
from converter import kind
from office_suite import Organize, Pages, Sketch, Tables
from watchit import WatchIT


class UsabilityTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def tearDown(self):
        for widget in self.app.topLevelWidgets():
            widget.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)

    def wait_for(self, condition):
        deadline = time.monotonic() + 5
        while not condition() and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(.001)
        self.assertTrue(condition())

    def test_brand_hover_animates_without_moving_buttons_and_respects_no_motion(self):
        from ui import palette
        from PySide6.QtCore import QAbstractAnimation
        from PySide6.QtWidgets import QStyle
        window = Window()
        self.assertEqual(palette()['accent'], '#ffdd22')
        self.assertNotIn('#087e75', window.styleSheet())
        button = window.add_button
        geometry = button.geometry()
        motion_style = SimpleNamespace(styleHint=lambda hint, *_: 160 if hint == QStyle.SH_Widget_Animation_Duration else 1)
        with patch.object(button, 'style', return_value=motion_style):
            QApplication.sendEvent(button, QEvent(QEvent.Enter))
        animation = button._hover_animation
        effect = animation.targetObject()
        if not button._hover_frame.isHidden():
            self.assertEqual(animation.endValue(), 1)
            self.assertLessEqual(animation.duration(), 160)
            animation.setCurrentTime(animation.duration() // 2)
            self.assertGreater(effect.opacity(), 0)
            self.assertLess(effect.opacity(), 1)
            animation.setCurrentTime(animation.duration())
            self.assertEqual(effect.opacity(), 1)
        self.assertIsNone(button.graphicsEffect())
        self.assertEqual(button.geometry(), geometry)
        with patch.object(button, 'style', return_value=motion_style):
            QApplication.sendEvent(button, QEvent(QEvent.Leave))
        animation.setCurrentTime(animation.duration())
        self.assertTrue(button._hover_frame.isHidden())
        with patch.object(button, 'style', return_value=SimpleNamespace(styleHint=lambda *_: 0)):
            QApplication.sendEvent(button, QEvent(QEvent.Enter))
            self.assertEqual(animation.state(), QAbstractAnimation.Stopped)
            self.assertTrue(button._hover_frame.isHidden())
        button.setEnabled(False)
        QApplication.sendEvent(button, QEvent(QEvent.Enter))
        self.assertTrue(button._hover_frame.isHidden())
        window.close()

    def test_hover_clears_between_buttons_blank_space_and_window_deactivation(self):
        from PySide6.QtCore import QPoint
        from PySide6.QtTest import QTest
        from PySide6.QtWidgets import QWidget
        from ui import editor_theme
        window = QWidget()
        window.setStyleSheet(editor_theme())
        window.resize(400, 200)
        buttons = [QPushButton(label, window) for label in ('First', 'Second')]
        for index, button in enumerate(buttons):
            button.setGeometry(10 + index * 100, 10, 80, 40)
        window.show()
        QTest.qWait(30)
        for index, button in enumerate(buttons):
            QTest.mouseMove(window, QPoint(40 + index * 100, 25))
            QTest.qWait(180)
            self.assertTrue(button._hover_frame.isVisible())
            if index:
                self.assertTrue(buttons[0]._hover_frame.isHidden())
        QTest.mouseMove(window, QPoint(300, 150))
        self.app.processEvents()
        self.assertTrue(all(button._hover_frame.isHidden() for button in buttons))
        QTest.mouseMove(window, QPoint(40, 25))
        QTest.qWait(180)
        QApplication.sendEvent(window, QEvent(QEvent.WindowDeactivate))
        self.assertTrue(buttons[0]._hover_frame.isHidden())
        # A popup can take the pointer without delivering Leave to its old button.
        QApplication.sendEvent(buttons[0], QEvent(QEvent.Enter))
        QApplication.sendEvent(buttons[1], QEvent(QEvent.Enter))
        self.assertTrue(buttons[0]._hover_frame.isHidden())
        self.assertEqual(buttons[0]._hover_animation.targetObject().opacity(), 0)
        window.close()

    def test_qt_created_toolbar_button_keeps_hover_state_until_destroyed(self):
        import gc
        import weakref
        from PySide6.QtWidgets import QToolBar
        from ui import palette
        palette()
        from PySide6.QtCore import QPoint
        from PySide6.QtTest import QTest
        toolbar = QToolBar()
        toolbar.resize(200, 60)
        action = toolbar.addAction('Open')
        toolbar.show()
        QTest.qWait(20)
        QTest.mouseMove(toolbar, toolbar.actionGeometry(action).center())
        QTest.qWait(20)
        button = toolbar.widgetForAction(action)
        self.assertTrue(hasattr(button, '_hover_frame'), 'Qt-created button must retain its hover state')
        tracked = weakref.ref(button)
        del button
        gc.collect()
        self.assertIsNotNone(tracked(), 'Qt-created button wrapper must survive between hover events')
        button = toolbar.widgetForAction(action)
        QApplication.sendEvent(button, QEvent(QEvent.Leave))
        self.assertTrue(button._hover_frame.isHidden())
        self.assertEqual(button._hover_animation.targetObject().opacity(), 0)
        self.assertIn(button, self.app._hover_feedback._buttons)
        toolbar.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        self.assertNotIn(button, self.app._hover_feedback._buttons)

    def test_task_tab_and_format_hover_follows_pointer_without_selecting(self):
        from PySide6.QtCore import QPoint
        from PySide6.QtTest import QTest
        from PySide6.QtWidgets import QWidget, QTabBar, QComboBox, QStyle
        from ui import editor_theme
        window = QWidget()
        window.setStyleSheet(editor_theme())
        window.resize(900, 160)
        tabs = QTabBar(window)
        tabs.setExpanding(False)
        for label in ('Convert', 'Reduce size', 'Create ZIP', 'Extract ZIP'):
            tabs.addTab(label)
        tabs.setGeometry(0, 0, 850, 40)
        combo = QComboBox(window)
        combo.addItems(['PNG', 'JPEG'])
        combo.setGeometry(10, 60, 120, 35)
        window.show()
        QTest.qWait(20)
        motion = SimpleNamespace(styleHint=lambda hint, *_: 160 if hint == QStyle.SH_Widget_Animation_Duration else 1)
        with patch.object(tabs, 'style', return_value=motion), patch.object(combo, 'style', return_value=motion):
            for index in range(tabs.count()):
                QTest.mouseMove(tabs, tabs.tabRect(index).center())
                self.app.processEvents()
                self.assertEqual(tabs._hover_frame.geometry(), tabs.tabRect(index))
                self.assertFalse(tabs._hover_frame.isHidden())
                self.assertEqual(tabs._hover_animation.endValue(), 1)
                self.assertEqual(tabs.currentIndex(), 0, 'Hover must not change the task')
            QTest.mouseMove(combo, combo.rect().center())
            self.app.processEvents()
            self.assertTrue(tabs._hover_frame.isHidden())
            self.assertFalse(combo._hover_frame.isHidden())
            self.assertEqual(combo._hover_frame.geometry(), combo.rect())
            QTest.mouseClick(combo, Qt.LeftButton, pos=combo.rect().center())
            self.assertTrue(combo.view().isVisible(), 'Highlight must not intercept the dropdown click')
            combo.hidePopup()
            QTest.mouseMove(window, QPoint(880, 140))
            self.app.processEvents()
            self.assertTrue(combo._hover_frame.isHidden())
            tabs.setTabEnabled(2, False)
            QTest.mouseMove(tabs, tabs.tabRect(2).center())
            self.app.processEvents()
            self.assertTrue(tabs._hover_frame.isHidden())
            self.assertEqual(tabs.cursor().shape(), Qt.ArrowCursor)
            QTest.mouseMove(tabs, QPoint(tabs.width() - 1, 20))
            self.assertTrue(tabs._hover_frame.isHidden())
        window.close()

    def test_dazzle_templates_preview_and_preserve_slide_colors(self):
        from office_templates import TEMPLATES, create_document, SLIDE_COLORS
        from office_suite import Dazzle
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            for name in TEMPLATES['Dazzle']:
                target = create_document('Dazzle', name, Path(folder)/name)
                original = target.read_bytes()
                editor = Dazzle()
                editor.open(target)
                self.assertTrue(all(colors == SLIDE_COLORS for colors in editor.slide_styles), name)
                self.assertIn('#ffdd22', editor.title.styleSheet())
                editor.title.setPlainText('Updated title')
                saved = Path(folder)/('Edited ' + name)
                editor._write(saved)
                self.assertEqual(editor._read_slide_colors(Presentation(saved).slides[0]), SLIDE_COLORS)
                self.assertEqual(target.read_bytes(), original)
                editor._write(Path(folder)/'colors.dazzle')
                editor.open(Path(folder)/'colors.dazzle')
                self.assertEqual(editor.slide_styles[0], SLIDE_COLORS)
                editor.add()
                self.assertEqual(len(editor.slides), len(editor.slide_styles))
                editor.delete()
                self.assertEqual(len(editor.slides), len(editor.slide_styles))
                editor.dirty = False
                editor.close()
            window = Window()
            dazzle = window._open_workspace('Dazzle')
            menu = next(action.menu() for action in dazzle.menuBar().actions() if action.text() == 'Templates')
            self.assertIn('Pitch deck', [action.text() for action in menu.actions()])
            pdf_path = Path(folder)/'themed.pdf'
            dazzle.title.setPlainText('Themed export')
            with patch('office_suite.pick_save', return_value=pdf_path):
                dazzle.export_pdf()
            import pypdfium2 as pdfium
            with pdfium.PdfDocument(pdf_path) as pdf:
                page = pdf[0]
                bitmap = page.render(scale=.1)
                try:
                    self.assertEqual(bitmap.to_pil().getpixel((2, 2))[:3], (245, 243, 235))
                finally:
                    bitmap.close()
                    page.close()
            import json
            bad = Path(folder)/'invalid.dazzle'
            bad.write_text(json.dumps({'slides':[['Bad', 'Text']], 'styles':[dict(SLIDE_COLORS, background='red; border:0')]}))
            before = list(dazzle.slides)
            dazzle.dirty = False
            with patch('office_suite.QMessageBox.warning') as warning:
                dazzle.open(bad)
                warning.assert_called_once()
            self.assertEqual(dazzle.slides, before)
            window.close()

    def test_windows_media_launch_waits_for_engine(self):
        with tempfile.TemporaryDirectory() as folder:
            media = Path(folder)/'video with spaces.mp4'
            media.write_bytes(b'fixture')
            window = Window()
            with patch.object(window, '_watchit') as launch:
                window.watchit = MagicMock(player=None)
                window.open_arguments(['--play', str(media)])
                launch.assert_called_once()
                self.assertEqual(window.pending_media, [([media.resolve()], False)])
                window.watchit.add_media.assert_not_called()
                window._watchit_ready(('vlc', 'instance', 'player'))
                window.watchit.add_media.assert_called_once_with([media.resolve()], autoplay=True)
                self.assertEqual(window.pending_media, [])
            window.watchit = None
            window.close()

    def test_template_opens_new_document_without_touching_existing_editor(self):
        with tempfile.TemporaryDirectory() as folder:
            window = Window()
            previous = Pages()
            previous.editor.setPlainText('Unsaved user work')
            window.workspaces['Pages'] = previous
            target = Path(folder)/'My report.docx'
            created = window.create_from_template('Pages', 'Report.docx', target)
            self.assertIsNot(created, previous)
            self.assertIn('Report title', created.editor.toPlainText())
            from PySide6.QtGui import QTextCursor
            cursor = QTextCursor(created.editor.document())
            cursor.movePosition(QTextCursor.Start)
            self.assertGreater(cursor.charFormat().fontPointSize(), 20)
            self.assertIsNone(created.imported_docx)
            created._write(target)
            from docx import Document
            self.assertEqual(Document(target).paragraphs[0].runs[0].font.size.pt, 28)
            self.assertEqual(previous.editor.toPlainText(), 'Unsaved user work')
            self.assertEqual(created.path, target)
            created.editor.document().setModified(False)
            previous.editor.document().setModified(False)
            window.close()

    def test_converter_layout_queue_and_background_preview(self):
        with tempfile.TemporaryDirectory() as folder:
            image = Path(folder) / "image.png"
            Image.new("RGB", (200, 100), "red").save(image)
            window = Window()
            window.show()
            for width, height in ((760, 540), (1080, 830), (1536, 864)):
                window.resize(width, height)
                self.app.processEvents()
                panel = window.findChild(QFrame, "controls")
                buttons = [button for button in panel.findChildren(QPushButton) if button.isVisible()]
                for button in buttons:
                    self.assertTrue(panel.rect().contains(button.geometry()), button.text())
                for index, button in enumerate(buttons):
                    for other in buttons[index + 1:]:
                        self.assertFalse(button.geometry().intersects(other.geometry()), (button.text(), other.text()))
            self.assertFalse(window.convert_button.isEnabled())
            for index, active in ((0, window.convert_button), (1, window.reduce_button),
                                  (2, window.compress_button), (3, window.zip_button)):
                window.task_tabs.setCurrentIndex(index)
                self.app.processEvents()
                self.assertTrue(active.isVisible())
            window.task_tabs.setCurrentIndex(0)
            window._add_files([image])
            self.wait_for(lambda: window.preview_worker is None)
            self.assertFalse(window.preview.pixmap().isNull())
            self.assertFalse(window.drop.isVisible())
            window._remove_selected()
            self.assertEqual(window.paths, [])
            self.assertFalse(window.convert_button.isEnabled())
            self.assertTrue(window.drop.isVisible())
            window.close()

    def test_batch_stays_locked_until_worker_finishes(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "note.txt"
            source.write_text("Temporary input", encoding="utf-8")
            window = Window()
            window.output_dir = Path(folder)
            window._add_files([source])
            release = threading.Event()
            with patch("app.convert", side_effect=lambda *_: release.wait(5)):
                window._convert()
                try:
                    self.assertFalse(window.convert_button.isEnabled())
                    self.assertFalse(window.clear_button.isEnabled())
                    self.assertFalse(window.destination.isEnabled())
                    self.assertFalse(window.close())
                finally:
                    release.set()
                    window.worker.wait(5000)
                    self.app.processEvents()
            self.assertTrue(window.convert_button.isEnabled())
            self.assertTrue(window.destination.isEnabled())
            window.close()

    def test_slow_preview_does_not_block_queue_and_stale_results_are_ignored(self):
        from PySide6.QtCore import QTimer
        from PySide6.QtGui import QImage
        with tempfile.TemporaryDirectory() as folder:
            picture = Path(folder) / "picture.png"
            text = Path(folder) / "note.txt"
            Image.new("RGB", (20, 20), "red").save(picture)
            text.write_text("Temporary text", encoding="utf-8")
            release = threading.Event()
            heartbeat = []
            window = Window()
            with patch("app.image_preview", side_effect=lambda _: release.wait(5)):
                window._add_files([picture, text])
                try:
                    QTimer.singleShot(0, lambda: heartbeat.append(True))
                    self.app.processEvents()
                    self.assertEqual(heartbeat, [True])
                    self.assertIsNotNone(window.preview_worker)
                    window.list.setCurrentRow(1)
                    window._preview_ready(picture.resolve(), QImage(10, 10, QImage.Format_RGB32))
                    self.assertIn("TXT", window.preview.text())
                finally:
                    release.set()
                    self.wait_for(lambda: window.preview_worker is None)
            window.close()

    def test_format_defaults_and_finished_batch_controls(self):
        self.assertEqual(kind(Path("sheet.xlsx")), "Spreadsheet")
        self.assertEqual(preferred_format(Path("movie.mp4")), "mkv")
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "note.txt"
            source.write_text("A local conversion", encoding="utf-8")
            window = Window()
            window.output_dir = Path(folder)
            window._add_files([source])
            window._convert()
            self.wait_for(lambda: not window.worker.isRunning())
            self.app.processEvents()
            self.assertTrue(window.convert_button.isEnabled())
            self.assertTrue((Path(folder) / "note.pdf").is_file())
            window.close()

    def test_watchit_fullscreen_and_playback_state(self):
        vlc = MagicMock()
        vlc.State = SimpleNamespace(Playing=1, Error=2, Ended=3)
        player = vlc.Instance.return_value.media_player_new.return_value
        player.audio_output_device_enum.return_value = None
        player.audio_output_device_get.return_value = None
        player.get_length.return_value = 1000
        player.get_time.return_value = 500
        player.get_state.return_value = vlc.State.Error
        with patch("watchit.vlc_runtime", return_value=vlc):
            window = WatchIT()
        window.showMaximized()
        self.assertEqual(window.play_button.width(), 32)
        self.assertEqual(window.play_button.iconSize().width(), 18)
        window.playlist_toggle.setChecked(True)
        self.app.processEvents()
        window.fullscreen()
        self.assertTrue(window.isFullScreen())
        self.assertFalse(window.playlist.isVisible())
        self.assertFalse(window.options_toggle.isVisible())
        window.exit_fullscreen()
        self.assertTrue(window.isMaximized())
        self.assertTrue(window.playlist.isVisible())
        window.mute.click()
        player.audio_set_mute.assert_called_with(True)
        self.assertEqual(window.mute.toolTip(), "Unmute (M)")
        window.index = 0
        window._refresh()
        self.assertIn("Could not play", window.statusBar().currentMessage())
        window.stop()
        self.assertEqual(window.seek.value(), 0)
        window.close()
        self.assertFalse(window.timer.isActive())
        window.show()
        self.assertTrue(window.timer.isActive())
        window.close()

    def test_watchit_engine_startup_leaves_converter_responsive(self):
        from PySide6.QtCore import QTimer
        release = threading.Event()
        heartbeat = []
        window = Window()
        with patch("watchit.create_playback_engine", side_effect=lambda: release.wait(5)), patch.object(window, "_watchit_ready"):
            window._watchit()
            try:
                self.assertIsNotNone(window.watchit)
                self.assertTrue(window.watchit.isVisible())
                self.assertIsNone(window.watchit.player)
                QTimer.singleShot(0, lambda: heartbeat.append(True))
                self.app.processEvents()
                self.assertEqual(heartbeat, [True])
                self.assertIsNotNone(window.watchit_loader)
                self.assertFalse(window.close())
            finally:
                release.set()
                self.wait_for(lambda: window.watchit_loader is None)
        window.close()

    def test_database_saves_preserve_hidden_records_and_detect_conflicts(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.sqlite"
            with closing(sqlite3.connect(path)) as db, db:
                db.execute("CREATE TABLE _freeconvert_meta (key TEXT PRIMARY KEY, value TEXT)")
                db.execute("INSERT INTO _freeconvert_meta VALUES ('owner', 'Organize')")
                db.execute("CREATE TABLE records (value TEXT, optional TEXT)")
                db.executemany("INSERT INTO records VALUES (?, NULL)", [(str(i),) for i in range(1005)])
                db.execute("CREATE TABLE other (value TEXT)")
            window = Organize()
            window.path = path
            window.refresh()
            window.tables.setCurrentRow(1)
            self.assertEqual(window.grid.rowCount(), 1000)
            window.grid.item(0, 0).setText("changed")
            with patch.object(QMessageBox, "question", return_value=QMessageBox.No):
                window.tables.setCurrentRow(0)
            self.assertEqual(window.tables.currentItem().text(), "records")
            window.grid.removeRow(1)
            window.add_record()
            window.grid.setItem(window.grid.rowCount() - 1, 0, QTableWidgetItem("new"))
            window.save()
            with closing(sqlite3.connect(path)) as db:
                self.assertEqual(db.execute("SELECT COUNT(*) FROM records").fetchone()[0], 1005)
                self.assertEqual(db.execute("SELECT optional FROM records WHERE value='changed'").fetchone(), (None,))
                self.assertEqual(db.execute("SELECT value FROM records WHERE value='1004'").fetchone(), ("1004",))
                db.execute("UPDATE records SET value='external' WHERE rowid=1")
                db.commit()
            window.grid.item(0, 0).setText("stale edit")
            with patch.object(QMessageBox, "warning") as warning:
                window.save()
                self.assertTrue(warning.called)
            with closing(sqlite3.connect(path)) as db:
                self.assertEqual(db.execute("SELECT value FROM records WHERE rowid=1").fetchone(), ("external",))
            window.dirty = False
            window.close()

    def test_editor_icons_and_formula_bar(self):
        from PySide6.QtWidgets import QToolBar
        from formulaic import Formulaic
        from office_suite import Dazzle
        for editor in (Pages(), Tables(), Dazzle(), Sketch(), Organize(), Formulaic()):
            if isinstance(editor, Dazzle):
                editor.show()
                self.app.processEvents()
                self.assertAlmostEqual(editor.slide_page.width() / editor.slide_page.height(), 16 / 9, places=2)
            for toolbar in editor.findChildren(QToolBar):
                self.assertEqual(toolbar.toolButtonStyle(), Qt.ToolButtonIconOnly)
                self.assertEqual(toolbar.iconSize().width(), 18)
                for item in toolbar.actions():
                    if item.property("iconName"):
                        self.assertFalse(item.icon().isNull(), item.text())
                        self.assertTrue(item.toolTip())
            editor.dirty = False
            editor.close()
        tables = Tables()
        tables.grid.setCurrentCell(2, 1)
        tables.formula.setText("=SUM(A1:A2)")
        tables._apply_formula()
        self.assertEqual(tables.grid.item(2, 1).text(), "=SUM(A1:A2)")
        self.assertEqual(tables.cell_address.text(), "B3")
        tables.dirty = False
        tables.close()

    def test_spreadsheet_headers_and_scaled_drawing(self):
        tables = Tables()
        tables.add_column()
        self.assertTrue(tables.dirty)
        self.assertEqual(tables.grid.horizontalHeaderItem(26).text(), "AA")
        tables.dirty = False
        tables.close()
        sketch = Sketch()
        sketch.canvas.resize(600, 400)
        self.assertEqual(sketch.canvas.image_rect().size().width(), 600)
        event = QMouseEvent(QMouseEvent.MouseButtonPress, QPointF(300, 200), QPointF(300, 200),
                            Qt.LeftButton, Qt.LeftButton, Qt.NoModifier)
        sketch.canvas.mousePressEvent(event)
        self.assertEqual(sketch.canvas.image.pixelColor(600, 400), sketch.canvas.color)
        sketch.dirty = False
        sketch.close()

    def test_new_word_document_retains_rich_text(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            editor = Pages()
            editor.editor.setHtml("<p><b>Bold</b> <i>italic</i> <u>underlined</u></p><ul><li>Bullet</li></ul>")
            path = Path(folder) / "formatted.docx"
            editor._write(path)
            doc = Document(path)
            self.assertTrue(doc.paragraphs[0].runs[0].bold)
            self.assertTrue(any(run.italic for run in doc.paragraphs[0].runs))
            self.assertTrue(any(run.underline for run in doc.paragraphs[0].runs))
            self.assertEqual(doc.paragraphs[1].style.name, "List Bullet")
            editor.close()


if __name__ == "__main__":
    unittest.main()
