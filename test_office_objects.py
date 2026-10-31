"""Round-trip checks for editable Office objects and nondestructive PDF organization."""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import io,tempfile,unittest,json
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from PIL import Image
from PySide6.QtWidgets import QApplication,QTableWidgetItem
from PySide6.QtCore import QCoreApplication,QEvent,Qt
from pptx import Presentation
from docx import Document
from openpyxl import Workbook,load_workbook
from office_suite import Dazzle,Pages,Tables
from office_insert import chart_data
from office_documents import insert_picture,insert_table,write_rich_docx
from office_spreadsheets import insert_image as sheet_image,insert_chart as sheet_chart,insert_video as sheet_video
from pdf_pages import page_range,write_pages,split_groups

class OfficeObjectsTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def tearDown(self):
        for widget in self.app.topLevelWidgets():widget.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)

    def test_open_action_uses_file_picker(self):
        deck=Dazzle()
        with patch('dazzle.QFileDialog.getOpenFileName',return_value=('', '')) as picker:
            next(a for a in deck.toolbar.actions() if a.text()=='Open').trigger()
            picker.assert_called_once()
    def picture(self,root):
        path=root/'picture.png';Image.new('RGB',(200,120),'red').save(path);return path
    def test_dazzle_native_objects_fonts_animation_project_and_undo(self):
        from dazzle import effects
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);image=self.picture(root);editor=Dazzle()
            text=editor.insert_text('A real text box');editor.format_object(name='Arial',size=32,bold=True,color='#ffdd22');editor.animate_object(text,.7)
            editor.insert_shape('Ellipse');editor.insert_image(image);editor.insert_chart('Column',chart_data('Month,Sales\nJan,12\nFeb,24'));editor.insert_table(2,3);editor.background_color('#102030')
            self.assertEqual(len(editor.deck.slides[0].shapes),7)
            before=len(editor.deck.slides[0].shapes);editor.undo();self.assertEqual(editor.slide_styles[0]['background'],'#f5f3eb');editor.redo();self.assertEqual(len(editor.deck.slides[0].shapes),before)
            target=root/'objects.pptx';editor._write(target)
            deck=Presentation(target);slide=deck.slides[0]
            self.assertTrue(any(s.has_chart for s in slide.shapes));self.assertTrue(any(s.has_table for s in slide.shapes));self.assertTrue(any(s.shape_type==13 for s in slide.shapes))
            textbox=next(s for s in slide.shapes if s.has_text_frame and s.text=='A real text box');font=textbox.text_frame.paragraphs[0].runs[0].font
            self.assertEqual(font.name,'Arial');self.assertEqual(font.size.pt,32);self.assertTrue(font.bold);self.assertEqual(effects(slide)[textbox.shape_id],.7)
            original=target.read_bytes();editor._write(root/'objects.dazzle');editor.dirty=False;editor.open(root/'objects.dazzle');self.assertEqual(len(editor.deck.slides[0].shapes),before)
            editor._write(root/'from-project.pptx');self.assertEqual(target.read_bytes(),original)
            editor.add();editor.title.setPlainText('Second');editor.move_slide(-1);self.assertEqual(editor.slides[0][0],'Second');editor.delete();self.assertEqual(len(editor.deck.slides),1)
    def test_canvas_hit_testing_selects_real_objects(self):
        from PySide6.QtCore import QPointF
        from PySide6.QtTest import QTest
        with tempfile.TemporaryDirectory() as folder:
            editor=Dazzle();shape=editor.insert_image(self.picture(Path(folder)));editor.show();self.app.processEvents();editor.scene.clearSelection()
            point=editor.slide_page.mapFromScene(QPointF(shape.left/12700+shape.width/25400,shape.top/12700+shape.height/25400))
            self.assertIsNotNone(editor.slide_page.itemAt(point))
            QTest.mouseClick(editor.slide_page.viewport(),Qt.LeftButton,pos=point);self.app.processEvents()
            self.assertEqual(editor.selected_shape().shape_id,shape.shape_id)
            editor.resize_object(editor.selected_shape(),240,160);self.assertEqual(editor.selected_shape().width/12700,240)
            editor.undo();self.assertNotEqual(editor.selected_shape().width/12700,240)

    def test_all_chart_types_can_be_inserted_previewed_and_edited(self):
        from dazzle import chart_info
        editor=Dazzle()
        for kind in ('Column','Bar','Line','Pie','Scatter'):
            data=chart_data('X,Series\n1,10\n2,20')
            shape=editor.insert_chart(kind,data);self.assertIsNotNone(chart_info(shape),kind)
            changed=chart_data('X,Series\n1,30\n2,40')
            with patch('dazzle.ask_chart',return_value=(kind,changed)):editor.edit_object(shape)
            self.assertEqual(list(shape.chart.series[0].values),[30,40],kind)

    def test_dazzle_video_is_embedded_and_source_preserved(self):
        import imageio_ffmpeg,subprocess
        from dazzle import movie_blob
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);video=root/'clip.mp4'
            subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(),'-y','-f','lavfi','-i','color=c=red:s=160x90:d=0.3','-c:v','libx264',str(video)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            editor=Dazzle();shape=editor.insert_video(video);self.assertEqual(movie_blob(shape),video.read_bytes());text=editor.insert_text('Animated over video');editor.animate_object(text);self.assertTrue(editor.deck.slides[0]._element.xpath('.//p:video'));ids=[n.get('id') for n in editor.deck.slides[0]._element.xpath('.//p:cTn')];self.assertEqual(len(ids),len(set(ids)));editor._write(root/'video.pptx')
            with ZipFile(root/'video.pptx') as archive:self.assertTrue(any(n.startswith('ppt/media/') and n.endswith('.mp4') for n in archive.namelist()))
    def test_pages_rich_image_table_and_font_survive_docx_reopen(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);editor=Pages();editor.editor.setPlainText('Rich document');cursor=editor.editor.textCursor();cursor.movePosition(__import__("PySide6.QtGui",fromlist=["QTextCursor"]).QTextCursor.MoveOperation.End);editor.editor.setTextCursor(cursor)
            insert_picture(editor,self.picture(root));cursor=editor.editor.textCursor();cursor.insertBlock();editor.editor.setTextCursor(cursor);table=insert_table(editor,2,2);table.cellAt(0,0).firstCursorPosition().insertText('Cell A')
            target=root/'rich.docx';editor._write(target);doc=Document(target);self.assertEqual(len(doc.inline_shapes),1);self.assertEqual(doc.tables[0].cell(0,0).text,'Cell A')
            editor.editor.document().setModified(False);editor.open(target);self.assertIn('data:image/',editor.editor.toHtml());self.assertIn('Cell A',editor.editor.toPlainText())
            original=target.read_bytes();editor._write(root/'unchanged.docx');self.assertEqual((root/'unchanged.docx').read_bytes(),original)
    def test_tables_appends_native_charts_images_without_dropping_original_parts(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'original.xlsx';book=Workbook();sheet=book.active;sheet.append(['Month','Sales']);sheet.append(['Jan',12]);sheet.append(['Feb',24]);book.create_sheet('Keep')['A1']='Original';book.save(source);book.close()
            editor=Tables();editor._read(source);editor.grid.setCurrentCell(4,3);sheet_image(editor,self.picture(root));sheet_chart(editor,'Bar',(1,1,3,2));output=root/'edited.xlsx';editor._write(output)
            book=load_workbook(output);self.assertEqual(book['Keep']['A1'].value,'Original');self.assertEqual(len(book.active._charts),1);self.assertEqual(len(book.active._images),1);book.close()
            reopened=Tables();reopened._read(output);self.assertEqual(len(reopened.inserted_objects),2)
            editor.grid.setItem(2,1,QTableWidgetItem('30'));editor._write(output);book=load_workbook(output);self.assertEqual(len(book.active._charts),1);self.assertEqual(book.active['B3'].value,30);book.close()
            # Append a second drawing to a workbook which already has a drawing tree.
            editor.grid.setCurrentCell(12,4);sheet_image(editor,self.picture(root));editor._write(output);book=load_workbook(output);self.assertEqual(len(book.active._images),2);self.assertEqual(len(book.active._charts),1);book.close()
            with ZipFile(source) as old,ZipFile(output) as new:self.assertEqual(old.read('xl/worksheets/sheet2.xml'),new.read('xl/worksheets/sheet2.xml'))
    def test_pdf_page_cards_reorder_remove_and_rotate_without_changing_input(self):
        from reportlab.pdfgen import canvas
        from pdf_pages import PdfPages
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'pages.pdf';pdf=canvas.Canvas(str(path))
            for number in range(3):pdf.drawString(20,700,str(number));pdf.showPage()
            pdf.save();original=path.read_bytes();dialog=PdfPages(path);dialog.timer.stop()
            dialog.pages.item(2).setSelected(True);dialog.move(-1)
            self.assertEqual([record[0] for record in dialog.records()],[1,3,2])
            dialog.rotate(90);self.assertEqual(dialog.records()[1],(3,90));dialog.remove();self.assertEqual([record[0] for record in dialog.records()],[1,2]);self.assertEqual(path.read_bytes(),original);dialog.close()

    def test_pdf_page_ranges_reorder_rotate_split_preserve_source(self):
        from reportlab.pdfgen import canvas
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'source.pdf';c=canvas.Canvas(str(source))
            for index in range(1,5):c.drawString(40,700,f'PAGE {index}');c.showPage()
            c.save();original=source.read_bytes();self.assertEqual(page_range('4-2,1',4),[4,3,2,1])
            write_pages(source,root/'arranged.pdf',[(4,90),(2,0)]);reader=PdfReader(root/'arranged.pdf');self.assertIn('PAGE 4',reader.pages[0].extract_text());self.assertEqual(reader.pages[0].rotation,90)
            split_groups(source,root/'parts',[[(1,0),(2,0)],[(3,180),(4,0)]])
            self.assertEqual(len(PdfReader(root/'parts/part-001.pdf').pages),2);self.assertEqual(source.read_bytes(),original)
            for invalid in ['1,1','0','1-999999999','1,,2','-2']:
                with self.assertRaises(ValueError):page_range(invalid,4)
            with self.assertRaises(FileExistsError):write_pages(source,source,[(1,0)])

if __name__=='__main__':unittest.main()
