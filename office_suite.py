"""Drunk Peanuts's local document, table, slide, drawing, and database workspaces."""

from __future__ import annotations

import csv
import re
import shutil
import sqlite3
from contextlib import closing, contextmanager
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from native_office import advanced_content, csv_delimiter
from office_surgery import edit_xlsx_cells, replace_runs, replace_text_frame
from PySide6.QtCore import QSignalBlocker, Qt, Signal, QPoint, QRect, QSize, Slot, QEvent
from PySide6.QtGui import QAction, QColor, QFont, QFontDatabase, QImage, QPainter, QPen, QTextListFormat
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QColorDialog, QFileDialog, QHBoxLayout, QInputDialog, QLabel,
    QListWidget, QMainWindow, QMessageBox, QPushButton, QSplitter, QTableWidget,
    QTableWidgetItem, QTextEdit, QToolBar, QVBoxLayout, QWidget, QComboBox, QSpinBox, QLineEdit, QFrame)
from ui import action as icon_action, editor_theme, refresh_icons
from office_templates import SLIDE_COLORS, style_slide


from office_formulas import calculate_cell


from office_base import Workspace, button, pick_open, pick_save


class Pages(Workspace):
    """Rich text editing with native HTML, TXT, DOCX and PDF output."""

    def __init__(self):
        super().__init__("Pages")
        self.imported_docx = None
        self.original_paragraphs = []
        self.editor = QTextEdit()
        self.editor.setAcceptRichText(True)
        self.editor.setPlaceholderText("Start writing…")
        self.editor.setObjectName("paper")
        self.editor.setFont(QFont("Calibri", 12))
        self.editor.document().setDocumentMargin(56)
        self.editor.setMaximumWidth(850)
        area = QWidget()
        layout = QHBoxLayout(area)
        layout.setContentsMargins(28, 24, 28, 20)
        layout.addStretch(1)
        paper=QWidget();paper.setMaximumWidth(850);paper_layout=QVBoxLayout(paper);paper_layout.setContentsMargins(0,0,0,0);paper_layout.setSpacing(0)
        self.header_preview=QLabel();self.footer_preview=QLabel()
        for preview in (self.header_preview,self.footer_preview):
            preview.setTextFormat(Qt.PlainText);preview.setWordWrap(True);preview.setStyleSheet('background:white;color:#444;padding:10px 28px;font:12px Calibri;');preview.hide()
        paper_layout.addWidget(self.header_preview);paper_layout.addWidget(self.editor,1);paper_layout.addWidget(self.footer_preview)
        layout.addWidget(paper, 10)
        layout.addStretch(1)
        self.setCentralWidget(area)
        for label, method in [("New", self.new), ("Open", self.open), ("Save", self.save),
                              ("Save as", self.save_as), ("PDF", self.export_pdf),
                              ("Undo", self.editor.undo), ("Redo", self.editor.redo),
                              ("Cut", self.editor.cut), ("Copy", self.editor.copy), ("Paste", self.editor.paste)]:
            button(self.toolbar, label, method)
        self.addToolBarBreak()
        formatting = QToolBar("Formatting", self)
        formatting.setMovable(False)
        formatting.setToolButtonStyle(Qt.ToolButtonIconOnly)
        formatting.setIconSize(QSize(18, 18))
        self.addToolBar(formatting)
        self.font_family = QComboBox()
        self.font_family.addItems(QFontDatabase.families())
        self.font_family.setCurrentText("Calibri")
        self.font_family.setMaximumWidth(210)
        self.font_family.setToolTip("Font family")
        self.font_family.currentTextChanged.connect(self._set_editor_font)
        formatting.addWidget(self.font_family)
        self.font_size = QSpinBox()
        self.font_size.setRange(6, 96)
        self.font_size.setValue(12)
        self.font_size.setSuffix(" pt")
        self.font_size.setToolTip("Font size")
        self.font_size.valueChanged.connect(self._set_font_size)
        formatting.addWidget(self.font_size)
        formatting.addSeparator()
        self.format_actions = {}
        for label, method in [
                              ("Bold", lambda: self.editor.setFontWeight(QFont.Bold if self.editor.fontWeight() < QFont.Bold else QFont.Normal)),
                              ("Italic", lambda: self.editor.setFontItalic(not self.editor.fontItalic())),
                              ("Underline", lambda: self.editor.setFontUnderline(not self.editor.fontUnderline())),
                              ("Bullets", self.bullets),
                              ("Align left", lambda: self.editor.setAlignment(Qt.AlignLeft)),
                              ("Align center", lambda: self.editor.setAlignment(Qt.AlignHCenter)),
                              ("Align right", lambda: self.editor.setAlignment(Qt.AlignRight))]:
            item = button(formatting, label, method)
            if label in {"Bold", "Italic", "Underline", "Align left", "Align center", "Align right"}:
                item.setCheckable(True)
                self.format_actions[label] = item
        self.editor.currentCharFormatChanged.connect(self._format_state)
        self.editor.cursorPositionChanged.connect(self._format_state)
        self.editor.textChanged.connect(self._count)
        self._count()
        from office_documents import setup_pages
        setup_pages(self)
        from office_editing import setup_pages_tools
        setup_pages_tools(self)
        from office_headers import setup_header_tools
        setup_header_tools(self)
        from office_productivity import setup_productivity
        setup_productivity(self)
        self.editor.document().setModified(False)

    @Slot(str)
    def _set_editor_font(self, font):
        self.rich_changed=True
        self.editor.setFontFamily(font)

    @Slot(int)
    def _set_font_size(self, size):
        self.rich_changed=True
        self.editor.setFontPointSize(size)

    @Slot()
    def _format_state(self):
        alignment = self.editor.alignment()
        values = {"Bold": self.editor.fontWeight() >= QFont.Bold, "Italic": self.editor.fontItalic(),
                  "Underline": self.editor.fontUnderline(), "Align left": bool(alignment & Qt.AlignLeft),
                  "Align center": bool(alignment & Qt.AlignHCenter), "Align right": bool(alignment & Qt.AlignRight)}
        for label, checked in values.items():
            self.format_actions[label].setChecked(checked)
        with QSignalBlocker(self.font_family), QSignalBlocker(self.font_size):
            self.font_family.setCurrentText(self.editor.currentFont().family())
            size = self.editor.fontPointSize()
            if size > 0:
                self.font_size.setValue(round(size))

    def _count(self):
        self.statusBar().showMessage(f"{len(self.editor.toPlainText().split())} words")

    def bullets(self):
        self.editor.textCursor().createList(QTextListFormat.ListDisc)

    def new(self):
        if not self._can_close():
            return
        self.path = None
        self.imported_docx = None
        self.original_paragraphs = []
        self.editor.clear()
        self.page_settings=None
        self.header_footer=None;self.header_footer_changed=False;self.original_body_html=None
        self.header_preview.hide();self.footer_preview.hide()
        from office_productivity import load_properties
        load_properties(self)
        self.rich_changed=False
        self.editor.document().setModified(False)

    def _can_close(self):
        if not self.editor.document().isModified():
            return True
        answer = QMessageBox.question(self, "Unsaved changes", "Discard changes?", QMessageBox.Yes | QMessageBox.No)
        return answer == QMessageBox.Yes

    def closeEvent(self, event):
        event.accept() if self._can_close() else event.ignore()

    def open(self, path=None):
        path = path or pick_open(self, "Open document", "Documents (*.docx *.html *.htm *.txt)")
        if not path or not self._can_close():
            return
        def read():
            self.page_settings=None
            self.header_footer=None
            self.header_footer_changed=False
            self.header_preview.hide();self.footer_preview.hide()
            from office_productivity import load_properties
            load_properties(self)
            if path.suffix.lower() == ".docx":
                doc = Document(path)
                load_properties(self,doc)
                self.original_paragraphs = [p.text for p in doc.paragraphs]
                from office_documents import docx_html
                self.editor.setHtml(docx_html(doc))
                from office_headers import load_headers
                load_headers(self,doc)
                self.rich_changed=False
                self.docx_loaded_rich=True
                self.imported_docx = path
            elif path.suffix.lower() in {".html", ".htm"}:
                self.editor.setHtml(path.read_text(encoding="utf-8"))
                self.imported_docx = None
            else:
                self.editor.setPlainText(path.read_text(encoding="utf-8"))
                self.imported_docx = None
            # Imported DOCX may contain structures this editor cannot retain.
            self.path = None if path.suffix.lower() == ".docx" else path
            self.editor.document().setModified(False)
            self.original_body_html=self.editor.toHtml()
        self.report(read)

    def save(self):
        if self.path:
            self.report(lambda: self._write(self.path))
        else:
            self.save_as()

    def save_as(self):
        path = pick_save(self, "Save document", "Word document (*.docx);;HTML (*.html);;Plain text (*.txt)")
        if path:
            self.report(lambda: self._write(path))

    def _write(self, path):
        if self.pending_properties and path.suffix.lower()!='.docx':raise ValueError('Save as DOCX to retain pending document properties.')
        if self.header_footer and path.suffix.lower() in ('.html','.htm','.txt'):raise ValueError('Save as DOCX or export PDF to retain header/footer and page numbers.')
        if path.suffix.lower()=='.docx' and self.imported_docx and self.docx_loaded_rich and not self.editor.document().isModified() and not self.rich_changed:
            if path.resolve()==self.imported_docx.resolve() and self.path is None:raise ValueError('Save an imported Word document to a new file')
            if path.resolve()!=self.imported_docx.resolve():shutil.copyfile(self.imported_docx,path)
            self.path=self.imported_docx=path
            return
        if self.imported_docx and self.docx_loaded_rich and self.editor.document().isModified():self.rich_changed=True
        if path.suffix.lower() == ".docx" and self.rich_changed:
            from office_documents import write_rich_docx
            write_rich_docx(self,path)
            return
        if path.suffix.lower() == ".docx":
            lines = self.editor.toPlainText().split("\n")
            if self.imported_docx and path.resolve() == self.imported_docx.resolve() and self.path is None:
                raise ValueError("Save an imported Word document to a new file")
            if self.imported_docx and lines == self.original_paragraphs:
                if path.resolve() != self.imported_docx.resolve():
                    shutil.copyfile(self.imported_docx, path)
            elif self.imported_docx and len(lines) == len(self.original_paragraphs):
                doc = Document(self.imported_docx)
                for paragraph, old, new in zip(doc.paragraphs, self.original_paragraphs, lines):
                    if new != old:
                        replace_runs(paragraph, new)
                doc.save(path)
            else:
                if self.imported_docx:
                    doc = Document(self.imported_docx)
                    if doc.tables or advanced_content(self.imported_docx):
                        raise ValueError("Paragraph count changed in a complex Word file; save as HTML or TXT to avoid dropping content")
                doc = Document()
                for key,value in self.document_properties.items():setattr(doc.core_properties,key,value)
                block = self.editor.document().begin()
                while block.isValid():
                    paragraph = doc.add_paragraph(style="List Bullet" if block.textList() else None)
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    from docx.shared import Pt
                    alignment = block.blockFormat().alignment()
                    paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if alignment & Qt.AlignHCenter else
                                           WD_ALIGN_PARAGRAPH.RIGHT if alignment & Qt.AlignRight else WD_ALIGN_PARAGRAPH.LEFT)
                    iterator = block.begin()
                    while not iterator.atEnd():
                        fragment = iterator.fragment()
                        if fragment.isValid():
                            run = paragraph.add_run(fragment.text())
                            formatting = fragment.charFormat()
                            run.bold = formatting.fontWeight() >= QFont.Bold
                            run.italic = formatting.fontItalic()
                            run.underline = formatting.fontUnderline()
                            run.font.name = formatting.font().family() or self.font_family.currentText()
                            run.font.size = Pt(formatting.fontPointSize() if formatting.fontPointSize() > 0 else self.font_size.value())
                        iterator += 1
                    block = block.next()
                doc.save(path)
            self.imported_docx = path
            self.original_paragraphs = lines
        elif path.suffix.lower() in {".html", ".htm"}:
            path.write_text(self.editor.toHtml(), encoding="utf-8")
        elif path.suffix.lower() == ".txt":
            path.write_text(self.editor.toPlainText(), encoding="utf-8")
        else:
            raise ValueError("Choose DOCX, HTML, or TXT")
        self.path = path
        self.editor.document().setModified(False)

    def export_pdf(self):
        path = pick_save(self, "Export PDF", "PDF (*.pdf)")
        if path:
            def write():
                printer = QPrinter(QPrinter.HighResolution)
                printer.setOutputFormat(QPrinter.PdfFormat)
                printer.setOutputFileName(str(path))
                if self.page_settings:
                    from PySide6.QtGui import QPageSize,QPageLayout
                    from PySide6.QtCore import QMarginsF
                    settings=self.page_settings
                    printer.setPageSize(QPageSize(QPageSize.A4 if settings['paper']=='A4' else QPageSize.Letter))
                    printer.setPageOrientation(QPageLayout.Landscape if settings['landscape'] else QPageLayout.Portrait)
                    printer.setPageMargins(QMarginsF(*([settings['margin']]*4)),QPageLayout.Millimeter)
                self.editor.document().print_(printer)
                del printer  # Finalize Qt's embedded font resources before reading the PDF.
                if self.header_footer:
                    from office_headers import overlay_pdf
                    overlay_pdf(path,self.header_footer)
            self.report(write)


class Tables(Workspace):
    """Editable spreadsheet cells with CSV and XLSX persistence."""

    def __init__(self):
        super().__init__("Tables")
        self.imported_xlsx = None
        self.original_rows = []
        self.csv_delimiter = ","
        self.csv_encoding = "utf-8-sig"
        self.grid = QTableWidget(100, 26)
        self.grid.setHorizontalHeaderLabels([chr(65 + i) for i in range(26)])
        self.grid.setAlternatingRowColors(True)
        self.grid.verticalHeader().setDefaultSectionSize(30)
        self.grid.horizontalHeader().setDefaultSectionSize(112)
        area = QWidget()
        layout = QVBoxLayout(area)
        layout.setContentsMargins(12, 12, 12, 12)
        formula_row = QHBoxLayout()
        self.cell_address = QLineEdit("A1")
        self.cell_address.setReadOnly(True)
        self.cell_address.setMaximumWidth(80)
        self.cell_address.setAccessibleName("Selected cell")
        formula_row.addWidget(self.cell_address)
        mark = QLabel("fx")
        mark.setObjectName("formulaMark")
        formula_row.addWidget(mark)
        self.formula = QLineEdit()
        self.formula.setAccessibleName("Cell value or formula")
        self.formula.setPlaceholderText("Cell value or formula · press Enter to apply")
        self.formula.returnPressed.connect(self._apply_formula)
        formula_row.addWidget(self.formula, 1)
        layout.addLayout(formula_row)
        layout.addWidget(self.grid, 1)
        self.setCentralWidget(area)
        self.grid.currentCellChanged.connect(self._show_formula)
        self.grid.itemChanged.connect(self._show_formula)
        self.grid.itemChanged.connect(lambda _: setattr(self, "dirty", True))
        for label, method in [("New", self.new), ("Open", self.open), ("Save", self.save), ("Save as", self.save_as),
                              ("Add row", self.add_row),
                              ("Add column", self.add_column),
                              ("Calculate selected", self.calculate_selected)]:
            button(self.toolbar, label, method)
        self.statusBar().showMessage("CSV and XLSX · formulas: arithmetic, SUM, AVERAGE, MIN, MAX, COUNT")
        self.grid.setCurrentCell(0, 0)
        from office_spreadsheets import setup_tables
        setup_tables(self)
        from office_cells import setup_cell_tools
        setup_cell_tools(self)
        from office_print_layout import setup_print_tools
        setup_print_tools(self)
        from office_formulas import setup_formulas
        setup_formulas(self)

    def _show_formula(self, *_):
        row, column = self.grid.currentRow(), self.grid.currentColumn()
        if row < 0 or column < 0:
            return
        header = self.grid.horizontalHeaderItem(column)
        self.cell_address.setText(f"{header.text() if header else column + 1}{row + 1}")
        cell = self.grid.item(row, column)
        self.formula.setText(cell.text() if cell else "")

    def _apply_formula(self):
        row, column = self.grid.currentRow(), self.grid.currentColumn()
        if row >= 0 and column >= 0:
            item=self.grid.item(row,column)
            if item:item.setText(self.formula.text())
            else:self.grid.setItem(row, column, QTableWidgetItem(self.formula.text()))

    def add_row(self):
        self.grid.insertRow(self.grid.rowCount())
        self.dirty = True

    def add_column(self):
        self.grid.insertColumn(self.grid.columnCount())
        self._headers()
        self.dirty = True

    def _headers(self):
        names = []
        for index in range(self.grid.columnCount()):
            name = ""
            while index >= 0:
                index, remainder = divmod(index, 26)
                name = chr(65 + remainder) + name
                index -= 1
            names.append(name)
        self.grid.setHorizontalHeaderLabels(names)

    def calculate_selected(self):
        row, column = self.grid.currentRow(), self.grid.currentColumn()
        if row < 0 or column < 0:
            return
        name = ""
        index = column
        while index >= 0:
            index, letter = divmod(index, 26)
            name = chr(65 + letter) + name
            index -= 1
        address = f"{name}{row + 1}"
        try:
            result = calculate_cell(self._rows(), address)
        except (ValueError, SyntaxError, TypeError, ZeroDivisionError, OverflowError) as error:
            self.statusBar().showMessage(f"{address}: {error}")
            return
        self.statusBar().showMessage(f"{address} = {result:g}" if isinstance(result, (int, float)) else f"{address} = {result}")

    def new(self):
        if not self.discard_changes():
            return
        self.grid.clearContents()
        self.inserted_objects=[];self.object_list.clear();self.object_preview.clear()
        self.pending_styles={}
        from office_print_layout import load_layout
        load_layout(self)
        self.path = None
        self.imported_xlsx = None
        self.original_rows = []
        self.dirty = False

    def open(self, path=None):
        path = path or pick_open(self, "Open table", "Spreadsheets (*.xlsx *.csv)")
        if path and self.discard_changes():
            delimiter = None
            if path.suffix.lower() == ".csv":
                choice, accepted = QInputDialog.getItem(self, "CSV import", "Separator:",
                    ["Detect automatically", "Comma", "Semicolon", "Tab", "Pipe"], 0, False)
                if not accepted:
                    return
                delimiter = {"Comma": ",", "Semicolon": ";", "Tab": "\t", "Pipe": "|"}.get(choice)
            self.report(lambda: self._read(path, delimiter))

    def _read(self, path, delimiter=None):
        if path.suffix.lower() == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                rows = list(book.active.values)
            finally:
                book.close()
        elif path.suffix.lower() == ".csv":
            delimiter = delimiter or csv_delimiter(path)
            with path.open(newline="", encoding="utf-8-sig") as stream:
                rows = list(csv.reader(stream, delimiter=delimiter))
        else:
            raise ValueError("Choose XLSX or CSV")
        self.inserted_objects=[];self.object_list.clear();self.object_preview.clear()
        self.grid.setUpdatesEnabled(False)
        blocker = QSignalBlocker(self.grid)
        self.grid.clearContents()
        self.grid.setRowCount(max(100, len(rows)))
        self.grid.setColumnCount(max(26, max((len(row) for row in rows), default=0)))
        self._headers()
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                if value is not None:
                    self.grid.setItem(i, j, QTableWidgetItem(str(value)))
        del blocker
        self.grid.setUpdatesEnabled(True)
        self.imported_xlsx = path if path.suffix.lower() == ".xlsx" else None
        if path.suffix.lower() == ".csv":
            self.csv_delimiter = delimiter
        self.original_rows = [["" if value is None else str(value) for value in row] for row in rows]
        # Imported XLSX requires a new path, preserving the original file.
        self.path = None if path.suffix.lower() == ".xlsx" else path
        self.dirty = False
        if path.suffix.lower()==".xlsx":
            from office_spreadsheets import load_objects
            load_objects(self,path)
        from office_cells import load_cell_formats
        load_cell_formats(self,path)
        from office_print_layout import load_layout
        load_layout(self,path)
        from office_formulas import recalculate
        recalculate(self)

    def save(self):
        if self.path:
            self.report(lambda: self._write(self.path))
        else:
            self.save_as()

    def save_as(self):
        path = pick_save(self, "Save table", "Excel workbook (*.xlsx);;CSV (*.csv)")
        if path:
            delimiter, encoding = self.csv_delimiter, self.csv_encoding
            if path.suffix.lower() == ".csv":
                labels = {"Comma": ",", "Semicolon": ";", "Tab": "\t", "Pipe": "|"}
                choice, accepted = QInputDialog.getItem(self, "CSV export", "Separator:", list(labels),
                    list(labels.values()).index(delimiter) if delimiter in labels.values() else 0, False)
                if not accepted:
                    return
                delimiter = labels[choice]
                choice, accepted = QInputDialog.getItem(self, "CSV export", "Encoding:",
                    ["UTF-8 with BOM (Excel compatible)", "UTF-8"], 0 if encoding == "utf-8-sig" else 1, False)
                if not accepted:
                    return
                encoding = "utf-8-sig" if choice.startswith("UTF-8 with") else "utf-8"
            self.report(lambda: self._write(path, delimiter, encoding))

    def _rows(self):
        rows = [[self.grid.item(i, j).text() if self.grid.item(i, j) else ""
                 for j in range(self.grid.columnCount())] for i in range(self.grid.rowCount())]
        while rows and not any(rows[-1]):
            rows.pop()
        width = max((next((j + 1 for j in range(len(row) - 1, -1, -1) if row[j]), 0) for row in rows), default=0)
        return [row[:width] for row in rows]

    def _write(self, path, delimiter=None, encoding=None):
        rows = self._rows()
        if path.suffix.lower() == ".xlsx":
            if self.imported_xlsx and path.resolve() == self.imported_xlsx.resolve() and self.path is None:
                raise ValueError("Save an imported workbook to a new file")
            if self.imported_xlsx and rows == self.original_rows and not self.pending_styles and not self.pending_print_layout and not any(not obj.get("saved") for obj in self.inserted_objects):
                if path.resolve() != self.imported_xlsx.resolve():
                    shutil.copyfile(self.imported_xlsx, path)
                self.imported_xlsx = path
                self.path = path
                self.dirty = False
                return
            if self.imported_xlsx:
                edit_xlsx_cells(self.imported_xlsx, path, self.original_rows, rows)
            else:
                book = Workbook()
                try:
                    sheet = book.active
                    for i, row in enumerate(rows):
                        for j, value in enumerate(row):
                            sheet.cell(i + 1, j + 1).value = (float(value) if value and value.replace(".", "", 1).replace("-", "", 1).isdigit() else value)
                    book.save(path)
                finally:
                    book.close()
            from office_spreadsheets import save_objects
            save_objects(self,path,rows)
            from office_cells import save_cell_formats
            save_cell_formats(self,path)
            from office_print_layout import save_layout
            save_layout(self,path)
            self.imported_xlsx = path
            self.original_rows = rows
        elif path.suffix.lower() == ".csv":
            if self.pending_print_layout:raise ValueError('Save as XLSX to retain worksheet print settings.')
            if self.pending_styles:raise ValueError('Save as XLSX to retain cell formatting.')
            if any(not obj.get("saved") for obj in self.inserted_objects):raise ValueError("Save as XLSX to retain inserted charts, images and media links.")
            with path.open("w", newline="", encoding=encoding or self.csv_encoding) as stream:
                csv.writer(stream, delimiter=delimiter or self.csv_delimiter).writerows(rows)
            self.csv_delimiter = delimiter or self.csv_delimiter
            self.csv_encoding = encoding or self.csv_encoding
        else:
            raise ValueError("Choose XLSX or CSV")
        self.path = path
        self.dirty = False


class Canvas(QWidget):
    changed = Signal()

    def __init__(self):
        super().__init__()
        self.image = QImage(1200, 800, QImage.Format_ARGB32)
        self.image.fill(Qt.white)
        self.color = QColor("#102f39")
        self.pen_width = 4
        self.last = None
        self.eraser = False
        self.setMinimumSize(240, 160)

    def image_rect(self):
        size = self.image.size().scaled(self.size(), Qt.KeepAspectRatio)
        return QRect((self.width() - size.width()) // 2, (self.height() - size.height()) // 2, size.width(), size.height())

    def image_point(self, event):
        rect = self.image_rect()
        point = event.position().toPoint()
        if not rect.contains(point):
            return None
        return QPoint((point.x() - rect.x()) * self.image.width() // max(1, rect.width()),
                      (point.y() - rect.y()) * self.image.height() // max(1, rect.height()))

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.fillRect(self.rect(), QColor("#d9e5e7"))
        painter.drawImage(self.image_rect(), self.image)

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.last = self.image_point(event)
            if self.last is not None:
                self._stroke(self.last)

    def _stroke(self, point):
        painter = QPainter(self.image)
        painter.setRenderHint(QPainter.Antialiasing)
        painter.setPen(QPen(Qt.white if self.eraser else self.color, self.pen_width, Qt.SolidLine, Qt.RoundCap))
        if self.last == point:
            painter.drawPoint(point)
        else:
            painter.drawLine(self.last, point)
        painter.end()
        self.last = point
        self.update()
        self.changed.emit()

    def mouseMoveEvent(self, event):
        if self.last is not None and event.buttons() & Qt.LeftButton:
            point = self.image_point(event)
            if point is not None:
                self._stroke(point)

    def mouseReleaseEvent(self, event):
        self.last = None


class Sketch(Workspace):
    def __init__(self):
        super().__init__("Sketch")
        self.imported_image = None
        self.canvas = Canvas()
        self.setCentralWidget(self.canvas)
        self.canvas.changed.connect(lambda: setattr(self, "dirty", True))
        for label, method in [("New", self.new), ("Open", self.open), ("Save", self.save),
                              ("Save as", self.save_as), ("Pen color", self.color),
                              ("Eraser", self.eraser), ("Clear", self.clear)]:
            button(self.toolbar, label, method)
        self.toolbar.addSeparator()
        self.pen_size = QSpinBox()
        self.pen_size.setRange(1, 60)
        self.pen_size.setValue(self.canvas.pen_width)
        self.pen_size.setSuffix(" px")
        self.pen_size.setToolTip("Pen / eraser width")
        self.pen_size.valueChanged.connect(lambda width: setattr(self.canvas, "pen_width", width))
        self.toolbar.addWidget(self.pen_size)

    def new(self):
        if not self.discard_changes():
            return
        self.canvas.image = QImage(1200, 800, QImage.Format_ARGB32)
        self.imported_image = None
        self.canvas.image.fill(Qt.white)
        self.canvas.update()
        self.path = None
        self.dirty = False

    def clear(self):
        self.canvas.image.fill(Qt.white)
        self.canvas.update()
        self.dirty = True

    def color(self):
        color = QColorDialog.getColor(self.canvas.color, self)
        if color.isValid():
            self.canvas.color = color
            self.canvas.eraser = False

    def eraser(self):
        self.canvas.eraser = not self.canvas.eraser

    def open(self, path=None):
        path = path or pick_open(self, "Open image", "Images (*.png *.jpg *.jpeg *.bmp)")
        if path and self.discard_changes():
            image = QImage(str(path))
            if image.isNull():
                QMessageBox.warning(self, "Sketch", "Could not open image")
                return
            self.canvas.image = image.convertToFormat(QImage.Format_ARGB32)
            self.imported_image = path
            self.canvas.update()
            # Imported images may contain metadata the canvas cannot retain.
            self.path = None
            self.dirty = False

    def save(self):
        if self.path:
            self.report(lambda: self._write(self.path))
        else:
            self.save_as()

    def save_as(self):
        path = pick_save(self, "Save image", "PNG (*.png);;JPEG (*.jpg);;BMP (*.bmp)")
        if path:
            self.report(lambda: self._write(path))

    def _write(self, path):
        if not self.canvas.image.save(str(path)):
            raise OSError("Could not save image")
        self.path = path
        self.dirty = False


class Organize(Workspace):
    """SQLite database table browser and record editor."""

    def __init__(self):
        super().__init__("Organize")
        self.loaded_table = None
        self.original_records = {}
        self.records_editable = False
        self.tables = QListWidget()
        self.tables.setMinimumWidth(180)
        self.grid = QTableWidget()
        self.grid.setAlternatingRowColors(True)
        self.grid.verticalHeader().setDefaultSectionSize(30)
        self.grid.itemChanged.connect(lambda _: setattr(self, "dirty", True))
        split = QSplitter()
        split.addWidget(self.tables)
        split.addWidget(self.grid)
        split.setSizes([200, 900])
        self.setCentralWidget(split)
        self.tables.currentTextChanged.connect(self._select_table)
        for label, method in [("New database", self.new), ("Open database", self.open),
                              ("New table", self.new_table), ("Add record", self.add_record),
                              ("Delete record", self.delete_record), ("Save changes", self.save),
                              ("Import CSV", self.import_csv), ("Export CSV", self.export_csv)]:
            button(self.toolbar, label, method)

    @contextmanager
    def connect(self):
        if not self.path:
            raise ValueError("Create or open a database first")
        db = sqlite3.connect(self.path)
        try:
            with db:
                yield db
        finally:
            db.close()

    def new(self):
        path = pick_save(self, "Create database", "SQLite database (*.sqlite *.db)")
        if path and self.discard_changes():
            if path.exists():
                QMessageBox.warning(self, "Organize", "Choose a new filename; an existing database cannot be replaced here.")
                return
            with closing(sqlite3.connect(path)) as db:
                with db:
                    db.execute("CREATE TABLE IF NOT EXISTS _freeconvert_meta (key TEXT PRIMARY KEY, value TEXT)")
                    db.execute("INSERT OR REPLACE INTO _freeconvert_meta VALUES ('owner', 'Organize')")
            self.path = path
            self.refresh()
            self.dirty = False

    def open(self, path=None):
        path = path or pick_open(self, "Open database", "SQLite database (*.sqlite *.db)")
        if path and self.discard_changes():
            self.path = path
            self.report(self.refresh)
            self.dirty = False

    def refresh(self):
        with self.connect() as db:
            names = [x[0] for x in db.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' AND name != '_freeconvert_meta' ORDER BY name")]
        with QSignalBlocker(self.tables):
            self.tables.clear()
            self.tables.addItems(names)
        self.loaded_table = None
        self.original_records = {}
        self.grid.clear()
        self.grid.setRowCount(0)
        self.dirty = False
        if names:
            self.tables.setCurrentRow(0)

    def _select_table(self, name):
        if name != self.loaded_table and not self.discard_changes():
            with QSignalBlocker(self.tables):
                matches = self.tables.findItems(self.loaded_table or "", Qt.MatchExactly)
                if matches:
                    self.tables.setCurrentItem(matches[0])
            return
        self.show_table(name)

    def show_table(self, name):
        if not name:
            self.grid.clear()
            return
        def read():
            with self.connect() as db:
                cols = [x[1] for x in db.execute(f"PRAGMA table_info({quote(name)})")]
                editable = self._editable()
                # Organize tables use SQLite rowids; external tables stay read-only.
                identity = next((key for key in ("rowid", "_rowid_", "oid") if key not in {col.lower() for col in cols}), None)
                editable = editable and identity is not None
                definition = db.execute("SELECT sql FROM sqlite_master WHERE name=?", (name,)).fetchone()[0] or ""
                editable = editable and "WITHOUT ROWID" not in definition.upper()
                rows = db.execute(f"SELECT {identity + ', ' if editable else ''}* FROM {quote(name)} LIMIT 1001").fetchall()
            truncated = len(rows) > 1000
            rows = rows[:1000]
            self.rowid_name = identity
            self.records_editable = editable
            self.loaded_table = name
            self.original_records = {row[0]: tuple(row[1:]) for row in rows} if editable else {}
            self.grid.setUpdatesEnabled(False)
            self.grid.clear()
            self.grid.setRowCount(len(rows))
            self.grid.setColumnCount(len(cols))
            self.grid.setHorizontalHeaderLabels(cols)
            for i, row in enumerate(rows):
                for j, value in enumerate(row[1:] if editable else row):
                    item = QTableWidgetItem("" if value is None else str(value))
                    if editable:
                        item.setData(Qt.UserRole, row[0])
                    self.grid.setItem(i, j, item)
            self.grid.setUpdatesEnabled(True)
            self.grid.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed if editable else QAbstractItemView.NoEditTriggers)
            self.dirty = False
            self.statusBar().showMessage(f"{len(rows)} records" + (" · showing the first 1,000; other records are preserved" if truncated else "") + (" · read-only" if not editable else ""))
        self.report(read)

    def new_table(self):
        if not self.path:
            return
        if not self._editable():
            QMessageBox.warning(self, "Organize", "External databases are read-only. Create an Organize database to edit records.")
            return
        if not self.discard_changes():
            return
        name, ok = QInputDialog.getText(self, "New table", "Table name")
        if not ok or not name.strip():
            return
        columns, ok = QInputDialog.getText(self, "Columns", "Comma-separated column names")
        if not ok:
            return
        names = [x.strip() for x in columns.split(",") if x.strip()]
        if not names or len(set(names)) != len(names):
            QMessageBox.warning(self, "Organize", "Enter unique column names")
            return
        def create():
            with self.connect() as db:
                db.execute(f"CREATE TABLE {quote(name.strip())} ({', '.join(quote(x) + ' TEXT' for x in names)})")
            self.refresh()
            self.tables.setCurrentItem(self.tables.findItems(name.strip(), Qt.MatchExactly)[0])
        self.report(create)

    def add_record(self):
        if self.tables.currentItem() and self.records_editable and self._editable():
            self.grid.insertRow(self.grid.rowCount())
            self.dirty = True

    def delete_record(self):
        row = self.grid.currentRow()
        if row >= 0 and self.records_editable and self._editable():
            self.grid.removeRow(row)
            self.dirty = True

    def save(self):
        name = self.tables.currentItem().text() if self.tables.currentItem() else None
        if not name:
            return
        if not self.records_editable or not self._editable():
            QMessageBox.warning(self, "Organize", "External databases are read-only")
            return
        def write():
            with self.connect() as db:
                db.execute("BEGIN IMMEDIATE")
                columns = [self.grid.horizontalHeaderItem(j).text() for j in range(self.grid.columnCount())]
                quoted_table = quote(name)
                identity = quote(self.rowid_name)
                remaining = set(self.original_records)
                placeholders = ",".join("?" for _ in columns)
                for i in range(self.grid.rowCount()):
                    first = self.grid.item(i, 0)
                    record_id = first.data(Qt.UserRole) if first else None
                    values = [self.grid.item(i, j).text() if self.grid.item(i, j) else None for j in range(len(columns))]
                    if record_id in self.original_records:
                        original = self.original_records[record_id]
                        remaining.discard(record_id)
                        current = db.execute(f"SELECT * FROM {quoted_table} WHERE {identity}=?", (record_id,)).fetchone()
                        if current != original:
                            raise ValueError("The database changed since it was opened. Reopen it before saving.")
                        changed = [j for j, value in enumerate(values) if value != ("" if original[j] is None else str(original[j]))]
                        if changed:
                            assignments = ", ".join(f"{quote(columns[j])}=?" for j in changed)
                            db.execute(f"UPDATE {quoted_table} SET {assignments} WHERE {identity}=?", [values[j] for j in changed] + [record_id])
                    else:
                        db.execute(f"INSERT INTO {quoted_table} ({', '.join(map(quote, columns))}) VALUES ({placeholders})", values)
                for record_id in remaining:
                    current = db.execute(f"SELECT * FROM {quoted_table} WHERE {identity}=?", (record_id,)).fetchone()
                    if current != self.original_records[record_id]:
                        raise ValueError("The database changed since it was opened. Reopen it before saving.")
                    db.execute(f"DELETE FROM {quoted_table} WHERE {identity}=?", (record_id,))
            self.show_table(name)
            self.statusBar().showMessage("Changes saved")
            self.dirty = False
        self.report(write)

    def import_csv(self):
        if not self.path or not self._editable():
            return
        if not self.discard_changes():
            return
        path = pick_open(self, "Import CSV", "CSV (*.csv)")
        if not path:
            return
        def write():
            with path.open(newline="", encoding="utf-8-sig") as stream:
                rows = list(csv.reader(stream))
            if not rows or not rows[0] or len(set(rows[0])) != len(rows[0]):
                raise ValueError("CSV needs unique column headings")
            name = path.stem
            with self.connect() as db:
                db.execute(f"CREATE TABLE {quote(name)} ({', '.join(quote(x) + ' TEXT' for x in rows[0])})")
                for row in rows[1:]:
                    db.execute(f"INSERT INTO {quote(name)} VALUES ({','.join('?' for _ in rows[0])})", (row + [None] * len(rows[0]))[:len(rows[0])])
            self.refresh()
        self.report(write)

    def export_csv(self):
        if not self.tables.currentItem():
            return
        path = pick_save(self, "Export CSV", "CSV (*.csv)")
        if path:
            def write():
                with path.open("w", newline="", encoding="utf-8-sig") as stream:
                    out = csv.writer(stream)
                    out.writerow([self.grid.horizontalHeaderItem(j).text() for j in range(self.grid.columnCount())])
                    for i in range(self.grid.rowCount()):
                        out.writerow([self.grid.item(i, j).text() if self.grid.item(i, j) else "" for j in range(self.grid.columnCount())])
            self.report(write)

    def _editable(self):
        if not self.path:
            return False
        with self.connect() as db:
            try:
                return db.execute("SELECT value FROM _freeconvert_meta WHERE key='owner'").fetchone() == ("Organize",)
            except sqlite3.OperationalError:
                return False


def quote(identifier):
    if not identifier or "\x00" in identifier:
        raise ValueError("Invalid table or column name")
    return '"' + identifier.replace('"', '""') + '"'


from formulaic import Formulaic

from dazzle import Dazzle

WORKSPACES = {"Pages": Pages, "Tables": Tables, "Dazzle": Dazzle, "Sketch": Sketch, "Organize": Organize, "Formulaic": Formulaic}
