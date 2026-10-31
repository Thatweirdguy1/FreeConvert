"""Native speaker notes, canvas sizes and presentation controls must round trip."""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import tempfile,unittest
from pathlib import Path
from PySide6.QtCore import Qt,QEvent,QCoreApplication
from PySide6.QtWidgets import QApplication
from PySide6.QtTest import QTest
from pptx import Presentation
from dazzle import Dazzle,SlideShow
from office_editing import duplicate_slide
from office_ribbon import install_ribbon


class PresentingBatchTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def tearDown(self):
        for widget in self.app.topLevelWidgets():widget.deleteLater()
        QCoreApplication.sendPostedEvents(None,QEvent.DeferredDelete)

    def test_visible_default_slide_real_previews_and_preserved_dark_import(self):
        import base64
        from PySide6.QtGui import QImage
        from office_templates import style_slide
        editor=Dazzle()
        def preview():
            data=editor.thumbnail(0).split(',',1)[1]
            return QImage.fromData(base64.b64decode(data))
        original=editor.deck.slides[0]._element.xml
        self.assertEqual(preview().pixelColor(3,3).name(),'#f5f3eb')
        self.assertEqual(editor.deck.slides[0]._element.xml,original)
        self.assertFalse(editor.dirty)
        editor.insert_text('Visible on the new page')
        shape=editor.deck.slides[0].shapes[-1]
        self.assertEqual(str(shape.text_frame.paragraphs[0].font.color.rgb),'202B3D')
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'Existing dark deck.pptx'
            deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);style_slide(slide);deck.save(path)
            editor.dirty=False;editor.open(path)
            original=editor.deck.slides[0]._element.xml
            self.assertEqual(preview().pixelColor(3,3).name(),'#0b0b0b')
            self.assertEqual(editor.deck.slides[0]._element.xml,original)
            editor.add()
            self.assertEqual(editor.slide_styles[1]['background'],'#0b0b0b')
            editor.background_color('#9457ff')
            self.assertEqual(QImage.fromData(base64.b64decode(editor.thumbnail(1).split(',',1)[1])).pixelColor(3,3).name(),'#9457ff')

    def test_notes_size_duplication_undo_and_native_roundtrip(self):
        with tempfile.TemporaryDirectory() as folder:
            editor=Dazzle();editor.set_notes('Opening cues\nKeep this off the slide.');editor.add();editor.set_notes('Second slide')
            editor.select(0);duplicate_slide(editor)
            self.assertEqual(editor.deck.slides[2].notes_slide.notes_text_frame.text,'Opening cues\nKeep this off the slide.')
            editor.set_notes('Independent duplicate');editor.undo()
            self.assertEqual(editor.deck.slides[2].notes_slide.notes_text_frame.text,'Opening cues\nKeep this off the slide.')
            editor.redo();self.assertEqual(editor.deck.slides[0].notes_slide.notes_text_frame.text,'Opening cues\nKeep this off the slide.')
            before=[s.left for s in editor.deck.slides[0].shapes];editor.set_slide_size(10,7.5)
            self.assertEqual(before,[s.left for s in editor.deck.slides[0].shapes])
            for suffix in ('.pptx','.dazzle'):
                path=Path(folder)/('deck'+suffix);editor._write(path);loaded=Dazzle();loaded.open(path)
                self.assertEqual(loaded.deck.slide_width,9144000)
                self.assertEqual(loaded.deck.slides[1].notes_slide.notes_text_frame.text,'Second slide')
                self.assertEqual(loaded.deck.slides[2].notes_slide.notes_text_frame.text,'Independent duplicate')
                if suffix=='.pptx':self.assertEqual(Presentation(path).slides[0].notes_slide.notes_text_frame.text,'Opening cues\nKeep this off the slide.')
            with self.assertRaises(ValueError):editor.set_slide_size(float('nan'),7.5)
            with self.assertRaises(ValueError):editor.set_slide_size(0,7.5)

    def test_slideshow_navigation_blackout_and_ribbon_commands(self):
        editor=Dazzle();editor.add();editor.add();install_ribbon(editor)
        self.assertIn('Speaker notes',editor.ribbon_buttons);self.assertIn('Slide size',editor.ribbon_buttons)
        self.assertEqual(editor.ribbon_buttons['Present from beginning'].defaultAction().shortcut().toString(),'F5')
        show=SlideShow(editor,0);self.assertEqual(show.index,0)
        QTest.keyClick(show,Qt.Key_End);self.assertEqual(show.index,2)
        QTest.keyClick(show,Qt.Key_B);self.assertTrue(show.blackout)
        self.assertTrue(all(not item.isVisible() for item in show.scene().items()))
        QTest.keyClick(show,Qt.Key_Left);self.assertEqual(show.index,2)
        QTest.keyClick(show,Qt.Key_B);self.assertFalse(show.blackout)
        QTest.keyClick(show,Qt.Key_Home);self.assertEqual(show.index,0)
        QTest.keyClick(show,Qt.Key_Escape);self.assertFalse(show.isVisible())


if __name__=='__main__':unittest.main()
