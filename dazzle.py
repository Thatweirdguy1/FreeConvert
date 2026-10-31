"""Object-based PowerPoint authoring; python-pptx remains the document model."""
from __future__ import annotations
import base64
import copy
import html
import io
import json
import math
import mimetypes
import re
import shutil
import tempfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.oxml.xmlchemy import OxmlElement
from PySide6.QtCore import Qt, QRectF, QSignalBlocker, QVariantAnimation, QTimer
from PySide6.QtGui import QColor, QFont, QFontDatabase, QImage, QPainter, QPen, QTextDocument, QAction, QPixmap
from PySide6.QtWidgets import (QApplication, QGraphicsObject, QGraphicsScene, QGraphicsView, QDockWidget,
 QWidget, QVBoxLayout, QFormLayout, QDoubleSpinBox, QComboBox, QToolBar, QTextEdit, QPushButton,
 QColorDialog, QInputDialog, QFileDialog, QMessageBox, QListWidget, QAbstractItemView)
from office_insert import ask_chart, chart_png, image_bytes, video_poster, validate_chart
from office_templates import SLIDE_COLORS, style_slide
from office_surgery import replace_text_frame
from ui import icon, palette

from office_base import Workspace, button
from PySide6.QtWidgets import QHBoxLayout, QSplitter, QGraphicsItem
from PySide6.QtCore import QEvent

EMU=12700
NEW_SLIDE_COLORS=dict(background='#f5f3eb',title='#202b3d',body='#202b3d')
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
FC='urn:freeconvert:slide-animation:1'
CHARTS={'Column':XL_CHART_TYPE.COLUMN_CLUSTERED,'Bar':XL_CHART_TYPE.BAR_CLUSTERED,'Line':XL_CHART_TYPE.LINE,'Pie':XL_CHART_TYPE.PIE,'Scatter':XL_CHART_TYPE.XY_SCATTER}

def rgb(color, fallback='#eeeeee'):
    try:return '#'+str(color.rgb) if color.type==1 else fallback
    except (AttributeError,ValueError):return fallback

def font_of(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            if p.runs:return p.runs[0].font
        return shape.text_frame.paragraphs[0].font
    return None

def primary(slide):
    texts=[s for s in slide.shapes if s.has_text_frame]
    title=slide.shapes.title
    if title is None:title=texts[0] if texts else None
    body=next((s for s in texts if s.shape_id != title.shape_id and s.is_placeholder),None)
    if body is None:body=next((s for s in texts if s.shape_id != title.shape_id),None)
    return title,body

def movie_blob(shape):
    refs=shape._element.xpath('.//a:videoFile')
    return shape.part.related_part(refs[0].get('{'+NS['r']+'}link')).blob if refs else None

def effects(slide):
    result={}
    for element in slide._element.findall('.//{'+FC+'}fade'):
        try:
            shape_id=int(element.get('shape'));duration=float(element.get('seconds'))
            if shape_id>0 and math.isfinite(duration) and .1<=duration<=10:result[shape_id]=duration
        except (TypeError,ValueError):continue
    return result

def set_effect(slide, shape_id, seconds):
    previous=effects(slide);current=dict(previous)
    if seconds:current[shape_id]=seconds
    else:current.pop(shape_id,None)
    template=etree.fromstring((Path(__file__).resolve().parent/'assets/animation-fade.xml').read_bytes())
    timing=slide._element.find('p:timing',NS)
    extlst=slide._element.find('p:extLst',NS)
    if extlst is None:extlst=OxmlElement('p:extLst');slide._element.append(extlst)
    if timing is None:
        timing=copy.deepcopy(template)
        timing.find('.//p:cTn[@nodeType="mainSeq"]/p:childTnLst',NS).clear();timing.find('p:bldLst',NS).clear()
        slide._element.insert(list(slide._element).index(extlst),timing)
    seq=timing.find('.//p:cTn[@nodeType="mainSeq"]/p:childTnLst',NS)
    if seq is None:
        root=timing.find('.//p:cTn[@nodeType="tmRoot"]/p:childTnLst',NS)
        if root is None:raise ValueError('This imported animation tree cannot be edited safely.')
        sequence=copy.deepcopy(template.find('.//p:seq',NS));sequence.find('p:cTn/p:childTnLst',NS).clear();sequence.find('p:cTn',NS).set('id',str(max((int(n.get('id')) for n in timing.findall('.//p:cTn',NS) if n.get('id','').isdigit()),default=1)+1));root.append(sequence);seq=sequence.find('p:cTn/p:childTnLst',NS)
    for group in list(seq):
        targets={int(n.get('spid')) for n in group.findall('.//p:spTgt',NS) if n.get('spid','').isdigit()}
        if targets and targets.issubset(previous):seq.remove(group)
    build=timing.find('p:bldLst',NS)
    if build is None:build=OxmlElement('p:bldLst');timing.append(build)
    for node in list(build):
        if node.get('spid','').isdigit() and int(node.get('spid')) in previous:build.remove(node)
    for element in list(extlst):
        if element.get('uri')==FC:extlst.remove(element)
    extension=OxmlElement('p:ext');extension.set('uri',FC);extlst.append(extension)
    group=template.find('.//p:cTn[@nodeType="mainSeq"]/p:childTnLst',NS)[0]
    identifier=max((int(n.get('id')) for n in timing.findall('.//p:cTn',NS) if n.get('id','').isdigit()),default=2)+1
    for sid,duration in current.items():
        clone=copy.deepcopy(group)
        for node in clone.findall('.//p:cTn',NS):node.set('id',str(identifier));identifier+=1
        for node in clone.findall('.//p:spTgt',NS):node.set('spid',str(sid))
        for node in clone.findall('.//p:animEffect/p:cBhvr/p:cTn',NS):node.set('dur',str(round(duration*1000)))
        seq.append(clone);node=OxmlElement('p:bldP');node.set('spid',str(sid));node.set('grpId','0');build.append(node)
        node=etree.SubElement(extension,'{'+FC+'}fade');node.set('shape',str(sid));node.set('seconds',str(duration))
    # Native movie timing remains in the original root alongside the main sequence.

def native_chart_data(kind,data):
    validate_chart(data,kind)
    result=XyChartData() if kind=='Scatter' else CategoryChartData()
    if kind=='Scatter':
        for series in data['series']:
            target=result.add_series(series['name'])
            for x,y in zip(data['categories'],series['values']):target.add_data_point(float(x),y)
    else:
        result.categories=data['categories']
        for series in data['series']:result.add_series(series['name'],series['values'])
    return result

def shape_html(shape):
    blocks=[]
    for paragraph in shape.text_frame.paragraphs:
        runs=[]
        for run in paragraph.runs:
            f=run.font;pf=paragraph.font
            size=f.size or pf.size or Pt(20);family=f.name or pf.name or 'Calibri'
            css=f'font-family:{html.escape(family,quote=True)};font-size:{size.pt}px;color:{rgb(f.color,rgb(pf.color))};'
            if f.bold or (f.bold is None and pf.bold):css+='font-weight:bold;'
            if f.italic or (f.italic is None and pf.italic):css+='font-style:italic;'
            if f.underline:css+='text-decoration:underline;'
            runs.append('<span style="'+css+'">'+html.escape(run.text).replace('\v','<br>')+'</span>')
        align={2:'center',3:'right'}.get(paragraph.alignment,'left')
        blocks.append('<p style="margin:0;text-align:'+align+'">'+(''.join(runs) or '<br>')+'</p>')
    return ''.join(blocks)

def chart_info(shape):
    chart=shape.chart;kind=next((name for name,value in CHARTS.items() if chart.chart_type==value),'Column')
    if kind=='Scatter':categories=[float(n.text) for n in chart.series[0]._element.xpath('.//c:xVal/c:numRef/c:numCache/c:pt/c:v')]
    else:categories=[c.label for c in chart.plots[0].categories]
    if not categories:return None
    return kind,{'categories':categories,'series':[{'name':s.name,'values':list(s.values)} for s in chart.series]}

class SlideObject(QGraphicsObject):
    def __init__(self, owner, shape, interactive=True):
        super().__init__();self.owner=owner;self.native_shape=shape;self.interactive=interactive;self.preview_size=None;self.resizing=False
        self.setPos(shape.left/EMU,shape.top/EMU);self.setTransformOriginPoint(self.boundingRect().center());self.setRotation(shape.rotation)
        if interactive:self.setFlags(QGraphicsItem.ItemIsSelectable|QGraphicsItem.ItemIsMovable)
        self.image=None;self.text=None
        try:
            if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE,MSO_SHAPE_TYPE.MEDIA):
                self.image=QImage.fromData(shape.image.blob)
            elif shape.has_chart:
                data=chart_info(shape)
                if data:self.image=QImage.fromData(chart_png(data[1],data[0]))
        except (AttributeError,ValueError,KeyError):pass
        if shape.has_text_frame:
            self.text=QTextDocument();self.text.setDocumentMargin(4);self.text.setHtml(shape_html(shape));self.text.setTextWidth(max(1,shape.width/EMU))
    def boundingRect(self):
        w,h=self.preview_size or (self.native_shape.width/EMU,self.native_shape.height/EMU)
        return QRectF(0,0,max(1,w),max(1,h))
    def paint(self,painter,option,widget=None):
        rect=self.boundingRect();shape=self.native_shape;painter.setRenderHint(QPainter.Antialiasing)
        if self.image is not None and not self.image.isNull():painter.drawImage(rect,self.image)
        elif shape.has_table:
            table=shape.table;cw=rect.width()/len(table.columns);rh=rect.height()/len(table.rows)
            painter.setFont(QFont('Calibri',12))
            for row in range(len(table.rows)):
                for col in range(len(table.columns)):
                    cell=QRectF(col*cw,row*rh,cw,rh);painter.fillRect(cell,QColor('#333333' if row==0 else '#ffffff'));painter.setPen(QColor('#888888'));painter.drawRect(cell);painter.setPen(QColor('#ffdd22' if row==0 else '#111111'));painter.drawText(cell.adjusted(4,3,-4,-3),Qt.TextWordWrap,table.cell(row,col).text)
        else:
            fill=QColor(rgb(shape.fill.fore_color,'#39331a')) if hasattr(shape,'fill') and shape.fill.type==1 else QColor(Qt.transparent)
            line=QColor('#888888')
            try:line=QColor(rgb(shape.line.color,'#888888'))
            except (AttributeError,ValueError):pass
            painter.setPen(QPen(line,1));painter.setBrush(fill)
            auto=shape.auto_shape_type if shape.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE else None
            if auto==MSO_SHAPE.OVAL:painter.drawEllipse(rect)
            elif auto==MSO_SHAPE.ISOSCELES_TRIANGLE:
                from PySide6.QtGui import QPolygonF
                from PySide6.QtCore import QPointF
                painter.drawPolygon(QPolygonF([QPointF(rect.width()/2,0),rect.bottomRight(),rect.bottomLeft()]))
            elif auto in (MSO_SHAPE.RIGHT_ARROW,MSO_SHAPE.STAR_5_POINT):
                from PySide6.QtGui import QPolygonF
                from PySide6.QtCore import QPointF
                w,h=rect.width(),rect.height()
                if auto==MSO_SHAPE.RIGHT_ARROW:points=[(0,h*.25),(w*.65,h*.25),(w*.65,0),(w,h*.5),(w*.65,h),(w*.65,h*.75),(0,h*.75)]
                else:points=[(w/2+math.cos(-math.pi/2+i*math.pi/5)*w/2*(1 if i%2==0 else .45),h/2+math.sin(-math.pi/2+i*math.pi/5)*h/2*(1 if i%2==0 else .45)) for i in range(10)]
                painter.drawPolygon(QPolygonF([QPointF(x,y) for x,y in points]))
            elif shape.shape_type==MSO_SHAPE_TYPE.LINE:painter.drawLine(rect.topLeft(),rect.bottomRight())
            elif not shape.has_text_frame or auto:painter.drawRoundedRect(rect,4,4)
            if not shape.has_text_frame and not auto and shape.shape_type!=MSO_SHAPE_TYPE.LINE:
                painter.setPen(QColor('#eeeeee'));painter.drawText(rect,Qt.AlignCenter,'Object preserved\n'+shape.name)
        if self.text is not None:self.text.drawContents(painter,rect)
        if self.isSelected():
            pen=QPen(QColor('#9457ff'),1.5,Qt.DashLine);pen.setCosmetic(True)
            painter.setPen(pen);painter.setBrush(Qt.NoBrush);painter.drawRect(rect.adjusted(.5,.5,-.5,-.5));painter.fillRect(QRectF(rect.right()-8,rect.bottom()-8,8,8),QColor('#9457ff'))
    def mousePressEvent(self,event):
        self.before=self.pos()
        if self.interactive and self.isSelected() and QRectF(self.boundingRect().right()-12,self.boundingRect().bottom()-12,12,12).contains(event.pos()):
            self.resizing=True;self.setFlag(QGraphicsItem.ItemIsMovable,False);event.accept();return
        super().mousePressEvent(event)
    def mouseMoveEvent(self,event):
        if self.resizing:
            self.prepareGeometryChange();self.preview_size=(max(10,event.pos().x()),max(10,event.pos().y()));self.update();event.accept();return
        super().mouseMoveEvent(event)
    def mouseReleaseEvent(self,event):
        if self.resizing:
            self.resizing=False;w,h=self.preview_size or (self.native_shape.width/EMU,self.native_shape.height/EMU);self.owner.resize_object(self.native_shape,w,h);event.accept();return
        super().mouseReleaseEvent(event)
        if self.pos()!=self.before:self.owner.move_object(self.native_shape,self.pos().x(),self.pos().y())
    def mouseDoubleClickEvent(self,event):
        if self.native_shape.shape_type==MSO_SHAPE_TYPE.MEDIA:self.owner.play_video(self.native_shape)
        elif self.interactive:self.owner.edit_object(self.native_shape)
        else:super().mouseDoubleClickEvent(event)

class SlideView(QGraphicsView):
    def drawForeground(self,painter,rect):
        # Editor-only page edge: it stays visible on dark slides and never exports.
        pen=QPen(QColor('#898596'),1);pen.setCosmetic(True)
        painter.save();painter.setPen(pen);painter.setBrush(Qt.NoBrush)
        painter.drawRect(self.sceneRect().adjusted(.5,.5,-.5,-.5));painter.restore()
    def resizeEvent(self,event):
        super().resizeEvent(event)
        if self.scene():self.fitInView(self.sceneRect(),Qt.KeepAspectRatio)

class SlideShow(QGraphicsView):
    def __init__(self,owner,start=None):
        super().__init__();self.owner=owner;self.index=max(0,owner.index if start is None else start);self.setWindowTitle('Dazzle · Slideshow');self.setStyleSheet('background:black;border:none');self.animations=[];self.pending=[];self.blackout=False;self.show_slide();self.showFullScreen()
    def show_slide(self):
        for animation in self.animations:animation.stop()
        old=self.scene();self.animations=[];scene=self.owner.make_scene(self.index,False);self.setScene(scene);self.pending=[]
        if old:old.deleteLater()
        for item in reversed(scene.items()):
            seconds=effects(self.owner.deck.slides[self.index]).get(item.native_shape.shape_id)
            if seconds:item.setOpacity(0);self.pending.append((item,seconds))
        self.fitInView(scene.sceneRect(),Qt.KeepAspectRatio)
        if self.owner.deck.slides[self.index]._element.find('p:transition/p:fade',NS) is not None:
            for item in scene.items():
                if item.opacity()>0:
                    item.setOpacity(0);animation=QVariantAnimation(self);animation.setDuration(450);animation.setStartValue(0.0);animation.setEndValue(1.0);animation.valueChanged.connect(item.setOpacity);self.animations.append(animation);animation.start()
    def keyPressEvent(self,event):
        if event.key()==Qt.Key_Escape:self.close()
        elif event.key() in (Qt.Key_B,Qt.Key_Period):
            self.blackout=not self.blackout
            for item in self.scene().items():item.setVisible(not self.blackout)
            self.scene().setBackgroundBrush(QColor('black') if self.blackout else QColor(self.owner._read_slide_colors(self.owner.deck.slides[self.index])['background']))
        elif self.blackout:return
        elif event.key()==Qt.Key_Home:self.index=0;self.show_slide()
        elif event.key()==Qt.Key_End:self.index=len(self.owner.deck.slides)-1;self.show_slide()
        elif event.key() in (Qt.Key_Space,Qt.Key_Right,Qt.Key_PageDown):self.advance()
        elif event.key() in (Qt.Key_Left,Qt.Key_PageUp):self.index=max(0,self.index-1);self.show_slide()
        else:super().keyPressEvent(event)
    def mousePressEvent(self,event):
        if self.blackout:return
        item=self.itemAt(event.pos())
        if item is not None and hasattr(item,'native_shape') and item.native_shape.shape_type==MSO_SHAPE_TYPE.MEDIA:self.owner.play_video(item.native_shape);return
        if event.button()==Qt.LeftButton:self.advance()
    def advance(self):
        if self.pending:
            item,seconds=self.pending.pop(0);animation=QVariantAnimation(self);animation.setDuration(round(seconds*1000));animation.setStartValue(0.0);animation.setEndValue(1.0);animation.valueChanged.connect(item.setOpacity);self.animations.append(animation);animation.start()
        elif self.index+1<len(self.owner.deck.slides):self.index+=1;self.show_slide()
        else:self.close()
    def resizeEvent(self,event):
        super().resizeEvent(event)
        if self.scene():self.fitInView(self.sceneRect(),Qt.KeepAspectRatio)
    def closeEvent(self,event):
        for animation in self.animations:animation.stop()
        scene=self.scene();self.setScene(None)
        if scene:scene.deleteLater()
        super().closeEvent(event)


class Dazzle(Workspace):
    def __init__(self):
        super().__init__('Dazzle')
        self.ready=True;self.loading=False;self.history=[];self.future=[];self._native_changed=False;self._loaded_path=None;self.players=[];self.imported_pptx=None;self.original_slides=[];self.index=0
        self.preview_revision=0;self._thumbnail_cache={}
        self.title=QTextEdit(self);self.body=QTextEdit(self);self.title.hide();self.body.hide()
        self.list=QListWidget();self.list.setMaximumWidth(225)
        self.slide_page=SlideView();self.slide_page.setFrameShape(QGraphicsView.NoFrame)
        self.slide_stage=QWidget();self.slide_stage.setObjectName('slideStage');self.slide_stage.setStyleSheet('QWidget#slideStage {background:#25252f;}');layout=QHBoxLayout(self.slide_stage);layout.setContentsMargins(24,24,24,24);layout.addStretch();layout.addWidget(self.slide_page);layout.addStretch()
        split=QSplitter();split.addWidget(self.list);split.addWidget(self.slide_stage);split.setStretchFactor(1,1);self.setCentralWidget(split)
        self.deck=self._from_legacy([('','')],[dict(NEW_SLIDE_COLORS)]);self.scene=None
        for label,callback in [('New',self.new),('Open',self.open),('Save',self.save),('Save as',self.save_as),('PDF',self.export_pdf),('Add slide',self.add),('Delete slide',self.delete)]:button(self.toolbar,label,callback)
        self._tools();self.refresh();self.dirty=False
        self.title.textChanged.connect(self._primary_changed);self.body.textChanged.connect(self._primary_changed);self.list.currentRowChanged.connect(self.select);self.slide_stage.installEventFilter(self)
        from office_editing import setup_slide_tools
        setup_slide_tools(self)
        from office_editing import command
        command(self,'Slide','Speaker notes',self.edit_notes,'pen')
        command(self,'Design','Slide size',self.choose_slide_size,'settings')
        command(self,'Slide','Present from beginning',lambda:self.present(0),'play','F5')
    def eventFilter(self,watched,event):
        if watched is self.slide_stage and event.type()==QEvent.Resize:
            scale=min(max(1,watched.width()-48)/(self.deck.slide_width/EMU),max(1,watched.height()-48)/(self.deck.slide_height/EMU))
            self.slide_page.setFixedSize(round(self.deck.slide_width/EMU*scale),round(self.deck.slide_height/EMU*scale))
        return super().eventFilter(watched,event)
    def save(self):
        if self.path:self.report(lambda:self._write(self.path))
        else:self.save_as()
    def save_as(self):
        name,_=QFileDialog.getSaveFileName(self,'Save presentation','','PowerPoint (*.pptx);;Dazzle project (*.dazzle)')
        if name:self.report(lambda:self._write(Path(name)))
    @staticmethod
    def _read_slide_colors(slide):
        from pptx.enum.dml import MSO_FILL, MSO_COLOR_TYPE
        colors = dict(background='#ffffff', title='#202b3d', body='#202b3d')
        fill = slide.background.fill
        if fill.type == MSO_FILL.SOLID and fill.fore_color.type == MSO_COLOR_TYPE.RGB:
            colors['background'] = '#' + str(fill.fore_color.rgb).lower()
            if QColor(colors['background']).lightness() < 128:
                colors.update(title='#f2f2ed', body='#f2f2ed')
        shapes = [shape for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        for role, shape in zip(('title', 'body'), shapes):
            fonts = [font for paragraph in shape.text_frame.paragraphs
                     for font in [*(run.font for run in paragraph.runs), paragraph.font]]
            rgb = next((font.color.rgb for font in fonts if font.color.type == MSO_COLOR_TYPE.RGB), None)
            if rgb is not None:
                colors[role] = '#' + str(rgb).lower()
        return colors

    def _from_legacy(self,slides,styles):
        deck=Presentation();deck.slide_width=Inches(13.333333);deck.slide_height=Inches(7.5)
        for pair,colors in zip(slides,styles):
            slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=pair[0];slide.placeholders[1].text=pair[1];style_slide(slide,colors)
        return deck
    def _tools(self):
        self.insert_menu=self.menuBar().addMenu('Insert');self.design_menu=self.menuBar().addMenu('Design');self.slide_menu=self.menuBar().addMenu('Slide');self.arrange_menu=self.menuBar().addMenu('Arrange')
        insertbar=QToolBar('Insert and design',self);insertbar.setMovable(False);insertbar.setToolButtonStyle(Qt.ToolButtonIconOnly);from PySide6.QtCore import QSize
        insertbar.setIconSize(QSize(18,18));self.addToolBarBreak();self.addToolBar(insertbar)
        def tool(menu,label,name,callback,shortcut=None):
            a=QAction(icon(name,palette()['accent']),label,self);a.setProperty('iconName',name);a.setToolTip(label);a.triggered.connect(lambda:self.report(callback));menu.addAction(a);insertbar.addAction(a)
            if shortcut:a.setShortcut(shortcut)
            return a
        tool(self.insert_menu,'Text box','title',lambda:self.insert_text('Text box'))
        tool(self.insert_menu,'Image','import',self.insert_image)
        tool(self.insert_menu,'Video','play',self.insert_video)
        tool(self.insert_menu,'Shape','pen',self.ask_shape)
        tool(self.insert_menu,'Chart / graph','table',self.insert_chart)
        tool(self.insert_menu,'Table','row',self.insert_table)
        tool(self.design_menu,'Background color','color',self.background_color)
        tool(self.design_menu,'Background image','open',self.background_image)
        tool(self.design_menu,'Slide transition','timer',self.choose_transition)
        tool(self.slide_menu,'Slideshow','slide',self.present,'Shift+F5')
        tool(self.slide_menu,'Move slide up','previous',lambda:self.move_slide(-1))
        tool(self.slide_menu,'Move slide down','next',lambda:self.move_slide(1))
        tool(self.menus['Edit'],'Undo','undo',self.undo,'Ctrl+Z');tool(self.menus['Edit'],'Redo','redo',self.redo,'Ctrl+Y')
        tool(self.arrange_menu,'Delete object','delete',self.delete_object,'Delete')
        tool(self.arrange_menu,'Bring to front','next',lambda:self.reorder_object(True))
        tool(self.arrange_menu,'Send to back','previous',lambda:self.reorder_object(False))
        dock=QDockWidget('Object properties',self);dock.setObjectName('slideObjectProperties');panel=QWidget();form=QFormLayout(panel)
        self.object_list=QListWidget();self.object_list.setMaximumHeight(125);self.object_list.currentRowChanged.connect(self._object_chosen);form.addRow('Objects',self.object_list)
        self.object_text=QTextEdit();self.object_text.setMaximumHeight(100);self.object_text.setAccessibleName('Selected object text');form.addRow('Text',self.object_text)
        apply=QPushButton('Apply text');apply.clicked.connect(lambda:self.report(lambda:self.edit_object(text=self.object_text.toPlainText())));form.addRow(apply)
        self.font_family=QComboBox();self.font_family.addItems(QFontDatabase.families());self.font_family.setMaximumWidth(210);self.font_family.setAccessibleName('Object font');self.font_family.currentTextChanged.connect(lambda v:self._format_if_selected(name=v));form.addRow('Font',self.font_family)
        self.font_size=QDoubleSpinBox();self.font_size.setRange(6,180);self.font_size.setSuffix(' pt');self.font_size.editingFinished.connect(lambda:self._format_if_selected(size=self.font_size.value()));form.addRow('Size',self.font_size)
        for label,name,key in [('Bold','bold','bold'),('Italic','italic','italic'),('Underline','underline','underline')]:
            tool(self.design_menu,label,name,lambda k=key:self.toggle_font(k))
        for label,callback in [('Text color',lambda:self.choose_object_color(True)),('Fill color',lambda:self.choose_object_color(False))]:
            b=QPushButton(label);b.clicked.connect(lambda checked=False,cb=callback:self.report(cb));form.addRow(b)
        self.geometry={}
        for key in ('X','Y','Width','Height','Rotation'):
            spin=QDoubleSpinBox();spin.setRange(0,10000 if key!='Rotation' else 359);spin.setDecimals(1);spin.setAccessibleName('Object '+key);spin.editingFinished.connect(self.apply_geometry);self.geometry[key]=spin;form.addRow(key,spin)
        self.animation=QComboBox();self.animation.addItems(['None','Fade on click']);self.animation.currentIndexChanged.connect(self.apply_animation);form.addRow('Animation',self.animation)
        self.effect_duration=QDoubleSpinBox();self.effect_duration.setRange(.1,10);self.effect_duration.setValue(.6);self.effect_duration.setSuffix(' s');self.effect_duration.editingFinished.connect(self.apply_animation);form.addRow('Duration',self.effect_duration)
        from PySide6.QtWidgets import QScrollArea
        scroll=QScrollArea();scroll.setWidgetResizable(True);scroll.setWidget(panel);dock.setWidget(scroll);dock.setMinimumWidth(225);self.addDockWidget(Qt.RightDockWidgetArea,dock)
        self.statusBar().showMessage('Select or drag objects · double-click text/chart/video · F5 presents · Escape exits')
    def _bytes(self):
        stream=io.BytesIO();self.deck.save(stream);return stream.getvalue()
    def change(self,operation):
        self._store();before=self._bytes();index=self.index
        try:result=operation()
        except Exception:
            self.deck=Presentation(io.BytesIO(before));self.index=index;self.refresh();raise
        self.history.append((before,index));self.history=self.history[-15:]
        while len(self.history)>1 and sum(len(v[0]) for v in self.history)>128*1024*1024:self.history.pop(0)
        self.future=[];self._native_changed=True;self.dirty=True;self.refresh();return result
    def undo(self):
        if not self.history:return
        self.future.append((self._bytes(),self.index));data,self.index=self.history.pop();self.deck=Presentation(io.BytesIO(data));self._native_changed=True;self.dirty=True;self.refresh()
    def redo(self):
        if not self.future:return
        self.history.append((self._bytes(),self.index));data,self.index=self.future.pop();self.deck=Presentation(io.BytesIO(data));self._native_changed=True;self.dirty=True;self.refresh()
    def _primary_changed(self):
        if self.ready and not self.loading:self._store();self.refresh_scene()
    def _store(self):
        if self.loading or self.index<0:return
        slide=self.deck.slides[self.index]
        for shape,widget in zip(primary(slide),(self.title,self.body)):
            if shape is not None and shape.text!=widget.toPlainText():
                replace_text_frame(shape,widget.toPlainText());self._native_changed=True;self.dirty=True
        self.slides[self.index]=(self.title.toPlainText(),self.body.toPlainText())
        if self.list.item(self.index):self.list.item(self.index).setText(self.title.toPlainText() or f'Slide {self.index+1}')
    def select(self,index):
        if index<0 or index>=len(self.deck.slides):return
        if index!=self.index:self._store()
        self.index=index;self.refresh()
    def refresh(self):
        self.loading=True
        try:
            self.slides=[tuple(s.text if s is not None else '' for s in primary(slide)) for slide in self.deck.slides]
            self.slide_styles=[self._read_slide_colors(slide) for slide in self.deck.slides]
            with QSignalBlocker(self.list):
                self.list.clear();self.list.addItems([pair[0] or f'Slide {i+1}' for i,pair in enumerate(self.slides)]);self.list.setCurrentRow(self.index)
            colors=self.slide_styles[self.index]
            for role,widget,value in zip(('title','body'),(self.title,self.body),self.slides[self.index]):
                with QSignalBlocker(widget):widget.setPlainText(value);widget.setStyleSheet('color:'+colors[role])
            self.refresh_scene()
        finally:self.loading=False
    def make_scene(self,index,interactive=True):
        scene=QGraphicsScene(self);scene.setSceneRect(0,0,self.deck.slide_width/EMU,self.deck.slide_height/EMU);scene.setBackgroundBrush(QColor(self._read_slide_colors(self.deck.slides[index])['background']))
        for z,shape in enumerate(self.deck.slides[index].shapes):
            if not interactive:
                # python-pptx style getters can create XML; previews must be read-only.
                from pptx.shapes.shapetree import SlideShapeFactory
                shape=SlideShapeFactory(copy.deepcopy(shape._element),self.deck.slides[index].shapes)
            item=SlideObject(self,shape,interactive);item.setZValue(z);scene.addItem(item)
        return scene
    def thumbnail(self,index):
        if not 0<=index<len(self.deck.slides):return ''
        if index in self._thumbnail_cache:return self._thumbnail_cache[index]
        from PySide6.QtCore import QBuffer,QIODevice
        height=max(1,round(256*self.deck.slide_height/self.deck.slide_width))
        image=QImage(256,height,QImage.Format_ARGB32);image.fill(Qt.transparent)
        scene=self.make_scene(index,False);painter=QPainter(image)
        try:
            painter.setRenderHints(QPainter.Antialiasing|QPainter.SmoothPixmapTransform)
            scene.render(painter,QRectF(0,0,image.width(),image.height()),scene.sceneRect())
        finally:painter.end();scene.deleteLater()
        buffer=QBuffer();buffer.open(QIODevice.WriteOnly);image.save(buffer,'PNG')
        result='data:image/png;base64,'+base64.b64encode(bytes(buffer.data())).decode('ascii')
        self._thumbnail_cache[index]=result
        return result
    def refresh_scene(self):
        if not getattr(self,'ready',False) or not hasattr(self,'deck'):return
        self.preview_revision+=1;self._thumbnail_cache.clear()
        selected=self.selected_shape().shape_id if getattr(self,'scene',None) and self.selected_shape() is not None else None
        old=self.scene;self.scene=self.make_scene(self.index);self.slide_page.setScene(self.scene);self.slide_page.fitInView(self.scene.sceneRect(),Qt.KeepAspectRatio)
        self.scene.selectionChanged.connect(self.selection_changed)
        with QSignalBlocker(self.object_list):
            self.object_list.clear();self.object_list.addItems([s.name for s in self.deck.slides[self.index].shapes])
        if selected:
            for item in self.scene.items():
                if item.native_shape.shape_id==selected:item.setSelected(True);break
        if old:old.deleteLater()
    def selected_shape(self):
        selected=self.scene.selectedItems() if getattr(self,'scene',None) else []
        return selected[0].native_shape if selected else None
    def _object_chosen(self,index):
        if self.loading or not self.scene:return
        self.scene.clearSelection()
        for item in self.scene.items():
            if item.native_shape.shape_id == self.deck.slides[self.index].shapes[index].shape_id:item.setSelected(True);break
    def selection_changed(self):
        shape=self.selected_shape()
        if shape is None:return
        self.loading=True
        try:
            self.object_text.setPlainText(shape.text if shape.has_text_frame else '')
            self.object_text.setEnabled(shape.has_text_frame)
            font=font_of(shape)
            self.font_family.setEnabled(font is not None);self.font_size.setEnabled(font is not None)
            if font:
                self.font_family.setCurrentText(font.name or 'Calibri');self.font_size.setValue(font.size.pt if font.size else 20)
            for key,value in zip(self.geometry,(shape.left/EMU,shape.top/EMU,shape.width/EMU,shape.height/EMU,shape.rotation)):self.geometry[key].setValue(value)
            seconds=effects(self.deck.slides[self.index]).get(shape.shape_id)
            self.animation.setCurrentIndex(1 if seconds else 0)
            if seconds:self.effect_duration.setValue(seconds)
            for i,s in enumerate(self.deck.slides[self.index].shapes):
                if s.shape_id==shape.shape_id:
                    with QSignalBlocker(self.object_list):self.object_list.setCurrentRow(i)
        finally:self.loading=False
    def choose(self,shape):
        self.scene.clearSelection()
        for item in self.scene.items():
            if item.native_shape.shape_id==shape.shape_id:item.setSelected(True);break
    def new(self):
        if not self.discard_changes():return
        self.deck=self._from_legacy([('','')],[dict(NEW_SLIDE_COLORS)]);self.index=0;self.path=self.imported_pptx=self._loaded_path=None;self.history=[];self.future=[];self._native_changed=False;self.refresh();self.dirty=False
    def add(self):
        def add():
            colors=dict(self.slide_styles[self.index])
            slide=self.deck.slides.add_slide(self.deck.slide_layouts[1]);style_slide(slide,colors);self.index=len(self.deck.slides)-1
        self.change(add)
    def delete(self):
        if len(self.deck.slides)<=1:return
        def remove():
            sid=self.deck.slides._sldIdLst[self.index];self.deck.part.drop_rel(sid.rId);self.deck.slides._sldIdLst.remove(sid);self.index=min(self.index,len(self.deck.slides)-1)
        self.change(remove)
    def move_slide(self,direction):
        target=self.index+direction
        if not 0<=target<len(self.deck.slides):return
        def move():
            ids=self.deck.slides._sldIdLst;node=ids[self.index];ids.remove(node);ids.insert(target,node);self.index=target
        self.change(move)
    def insert_text(self,text='Text box'):
        def add():
            s=self.deck.slides[self.index].shapes.add_textbox(Pt(100),Pt(150),Pt(350),Pt(90));s.text=text
            for p in s.text_frame.paragraphs:p.font.name='Calibri';p.font.size=Pt(24);p.font.color.rgb=RGBColor.from_string(self.slide_styles[self.index]['body'][1:])
            return s
        shape=self.change(add);self.choose(shape);return shape
    def insert_shape(self,kind='Rectangle'):
        kinds={'Rectangle':MSO_SHAPE.RECTANGLE,'Rounded rectangle':MSO_SHAPE.ROUNDED_RECTANGLE,'Ellipse':MSO_SHAPE.OVAL,'Triangle':MSO_SHAPE.ISOSCELES_TRIANGLE,'Arrow':MSO_SHAPE.RIGHT_ARROW,'Star':MSO_SHAPE.STAR_5_POINT}
        if kind not in kinds:raise ValueError('Unsupported shape')
        def add():
            s=self.deck.slides[self.index].shapes.add_shape(kinds[kind],Pt(140),Pt(160),Pt(240),Pt(140));s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string('FFDD22');return s
        shape=self.change(add);self.choose(shape);return shape
    def ask_shape(self):
        kind,ok=QInputDialog.getItem(self,'Insert shape','Shape',['Rectangle','Rounded rectangle','Ellipse','Triangle','Arrow','Star'],0,False)
        if ok:self.insert_shape(kind)
    def insert_image(self,path=None):
        if not path:
            name,_=QFileDialog.getOpenFileName(self,'Insert image','','Images (*.png *.jpg *.jpeg *.bmp *.webp)');path=Path(name) if name else None
        if not path:return
        blob=image_bytes(path)
        shape=self.change(lambda:self.deck.slides[self.index].shapes.add_picture(io.BytesIO(blob),Pt(100),Pt(120),width=Pt(360)))
        self.choose(shape);return shape
    def insert_video(self,path=None):
        if not path:
            name,_=QFileDialog.getOpenFileName(self,'Embed video','','Video (*.mp4 *.mov *.m4v *.webm *.avi)');path=Path(name) if name else None
        if not path:return
        poster=video_poster(path);mime=mimetypes.guess_type(str(path))[0] or 'video/mp4'
        shape=self.change(lambda:self.deck.slides[self.index].shapes.add_movie(str(path),Pt(100),Pt(120),Pt(480),Pt(270),poster_frame_image=io.BytesIO(poster),mime_type=mime))
        self.choose(shape);return shape
    def insert_chart(self,kind=None,data=None):
        if kind is None:
            selected=ask_chart(self)
            if not selected:return
            kind,data=selected
        chartdata=native_chart_data(kind,data)
        shape=self.change(lambda:self.deck.slides[self.index].shapes.add_chart(CHARTS[kind],Pt(90),Pt(130),Pt(600),Pt(330),chartdata))
        self.choose(shape);return shape
    def insert_table(self,rows=None,columns=None):
        if rows is None:
            rows,ok=QInputDialog.getInt(self,'Insert table','Rows',3,1,50)
            if not ok:return
            columns,ok=QInputDialog.getInt(self,'Insert table','Columns',3,1,20)
            if not ok:return
        if not 1<=rows<=50 or not 1<=columns<=20:raise ValueError('Table size exceeds 50 rows or 20 columns.')
        shape=self.change(lambda:self.deck.slides[self.index].shapes.add_table(rows,columns,Pt(90),Pt(150),Pt(650),Pt(250)))
        self.choose(shape);return shape
    def edit_object(self,shape=None,text=None):
        shape=shape or self.selected_shape()
        if shape is None:return
        if shape.has_table:
            row,ok=QInputDialog.getInt(self,'Edit table cell','Row',1,1,len(shape.table.rows))
            if not ok:return
            col,ok=QInputDialog.getInt(self,'Edit table cell','Column',1,1,len(shape.table.columns))
            if not ok:return
            value,ok=QInputDialog.getMultiLineText(self,'Edit cell','Text',shape.table.cell(row-1,col-1).text)
            if ok:self.change(lambda:setattr(shape.table.cell(row-1,col-1),'text',value))
        elif shape.has_chart:
            data=chart_info(shape)
            if data:
                result=ask_chart(self,data[1],data[0])
                if result:
                    if result[0]!=data[0]:raise ValueError('Insert a new chart to change its type.')
                    cd=native_chart_data(result[0],result[1]);self.change(lambda:shape.chart.replace_data(cd))
        elif shape.has_text_frame:
            if text is None:text,ok=QInputDialog.getMultiLineText(self,'Edit text box','Text',shape.text)
            else:ok=True
            if ok:
                def update():
                    if len(shape.text_frame.paragraphs)==len(text.split('\n')):replace_text_frame(shape,text)
                    else:shape.text=text
                self.change(update)
    def _format_if_selected(self,**values):
        if not self.loading and self.selected_shape() is not None:self.report(lambda:self.format_object(**values))
    def format_object(self,**values):
        shape=self.selected_shape()
        if shape is None or not shape.has_text_frame:return
        def apply():
            for paragraph in shape.text_frame.paragraphs:
                for font in [paragraph.font]+[r.font for r in paragraph.runs]:
                    for key,value in values.items():
                        if key=='color':font.color.rgb=RGBColor.from_string(value.lstrip('#'))
                        else:setattr(font,key,Pt(value) if key=='size' else value)
        self.change(apply)
    def toggle_font(self,key):
        shape=self.selected_shape();font=font_of(shape) if shape else None
        if font:self.format_object(**{key:not getattr(font,key)})
    def choose_object_color(self,text):
        shape=self.selected_shape()
        if shape is None:return
        color=QColorDialog.getColor(QColor('#ffdd22'),self)
        if not color.isValid():return
        if text:self.format_object(color=color.name())
        else:
            def apply():shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color.name()[1:])
            self.change(apply)
    def background_color(self,color=None):
        if color is None:
            chosen=QColorDialog.getColor(QColor(self.slide_styles[self.index]['background']),self)
            if not chosen.isValid():return
            color=chosen.name()
        if not re.fullmatch('#[0-9A-Fa-f]{6}',color):raise ValueError('Choose a valid color.')
        def apply():
            fill=self.deck.slides[self.index].background.fill;fill.solid();fill.fore_color.rgb=RGBColor.from_string(color[1:])
        self.change(apply)
    def choose_transition(self):
        value,ok=QInputDialog.getItem(self,'Slide transition','Effect',['None','Fade'],0,False)
        if ok:self.set_transition(value)
    def set_transition(self,value):
        if value not in ('None','Fade'):raise ValueError('Choose None or Fade.')
        def apply():
            slide=self.deck.slides[self.index]._element
            for old in list(slide):
                if etree.QName(old).localname=='transition':slide.remove(old)
            if value=='Fade':
                node=OxmlElement('p:transition');node.set('spd','med');node.set('advClick','1');node.append(OxmlElement('p:fade'));index=next((i for i,n in enumerate(slide) if etree.QName(n).localname in {'timing','extLst'}),len(slide));slide.insert(index,node)
        self.change(apply)
    def background_image(self,path=None):
        if path is None:
            name,_=QFileDialog.getOpenFileName(self,'Slide background image','','Images (*.png *.jpg *.jpeg *.webp)');path=Path(name) if name else None
        if not path:return
        blob=image_bytes(path)
        def apply():
            shapes=self.deck.slides[self.index].shapes;shape=shapes.add_picture(io.BytesIO(blob),0,0,self.deck.slide_width,self.deck.slide_height);shapes._spTree.remove(shape._element);shapes._spTree.insert(2,shape._element)
        self.change(apply)
    def move_object(self,shape,x,y):
        self.change(lambda:(setattr(shape,'left',Pt(max(0,x))),setattr(shape,'top',Pt(max(0,y)))))
    def resize_object(self,shape,width,height):
        if not all(math.isfinite(v) and v>0 for v in (width,height)):raise ValueError('Object dimensions must be positive.')
        self.change(lambda:(setattr(shape,'width',Pt(width)),setattr(shape,'height',Pt(height))))
    def apply_geometry(self):
        if self.loading:return
        shape=self.selected_shape()
        if shape is None:return
        values={key:spin.value() for key,spin in self.geometry.items()}
        if values['Width']<=0 or values['Height']<=0:return
        def apply():
            for key,name in [('X','left'),('Y','top'),('Width','width'),('Height','height')]:setattr(shape,name,Pt(values[key]))
            shape.rotation=values['Rotation']
        self.report(lambda:self.change(apply))
    def apply_animation(self,*_):
        if self.loading:return
        shape=self.selected_shape()
        if shape:self.report(lambda:self.change(lambda:set_effect(self.deck.slides[self.index],shape.shape_id,self.effect_duration.value() if self.animation.currentIndex() else None)))
    def animate_object(self,shape,seconds=.6):
        if shape is None:raise ValueError('Select an object first.')
        if seconds is not None and not .1<=seconds<=10:raise ValueError('Animation duration must be 0.1–10 seconds.')
        self.change(lambda:set_effect(self.deck.slides[self.index],shape.shape_id,seconds))
    def delete_object(self):
        shape=self.selected_shape()
        if shape:
            def remove():
                if shape.shape_id in effects(self.deck.slides[self.index]):set_effect(self.deck.slides[self.index],shape.shape_id,None)
                shape._element.getparent().remove(shape._element)
            self.change(remove)
    def reorder_object(self,front):
        shape=self.selected_shape()
        if shape:
            def move():
                tree=shape._element.getparent();tree.remove(shape._element);tree.insert(len(tree) if front else 2,shape._element)
            self.change(move)
    def present(self,start=None):
        self._store();self.show_window=SlideShow(self,start)

    def set_notes(self,text):
        if not isinstance(text,str) or len(text)>100000:raise ValueError('Notes must contain at most 100,000 characters.')
        def apply():
            frame=self.deck.slides[self.index].notes_slide.notes_text_frame
            if frame is None:raise ValueError('This notes master has no editable notes body.')
            frame.text=text
        self.change(apply)

    def edit_notes(self):
        from PySide6.QtWidgets import QDialog,QDialogButtonBox,QLabel
        slide=self.deck.slides[self.index]
        frame=slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        dialog=QDialog(self);dialog.setWindowTitle('Speaker notes');dialog.resize(560,370);layout=QVBoxLayout(dialog)
        layout.addWidget(QLabel('Private speaker notes saved in PPTX/Dazzle. Not shown on slides or PDF.\nApplying replaces this slide’s notes text and its text formatting.'))
        editor=QTextEdit();editor.setAcceptRichText(False);editor.setAccessibleName('Speaker notes text');editor.setPlainText(frame.text if frame is not None else '');layout.addWidget(editor)
        buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);layout.addWidget(buttons)
        if dialog.exec()==QDialog.Accepted:self.set_notes(editor.toPlainText())

    def set_slide_size(self,width,height):
        if not all(isinstance(v,(int,float)) and math.isfinite(v) and 1<=v<=56 for v in (width,height)):
            raise ValueError('Slide dimensions must be between 1 and 56 inches.')
        def apply():self.deck.slide_width=Inches(width);self.deck.slide_height=Inches(height)
        self.change(apply)
        # Trigger the existing stage fitting path after the native canvas changes.
        self.eventFilter(self.slide_stage,QEvent(QEvent.Resize))

    def choose_slide_size(self):
        from PySide6.QtWidgets import QDialog,QDialogButtonBox,QLabel
        dialog=QDialog(self);dialog.setWindowTitle('Slide size');layout=QFormLayout(dialog)
        preset=QComboBox();preset.addItems(['Custom','Widescreen (16:9)','Standard (4:3)','Portrait']);layout.addRow('Preset',preset)
        width=QDoubleSpinBox();height=QDoubleSpinBox()
        for spin,value,label in ((width,self.deck.slide_width/Inches(1),'Width'),(height,self.deck.slide_height/Inches(1),'Height')):
            spin.setRange(1,56);spin.setDecimals(3);spin.setValue(value);spin.setSuffix(' in');spin.setAccessibleName('Slide '+label.lower());layout.addRow(label,spin)
        def choose(index):
            if index:width.setValue([0,13.333,10,7.5][index]);height.setValue([0,7.5,7.5,10][index])
        preset.currentIndexChanged.connect(choose)
        layout.addRow(QLabel('Applies to the whole deck. Objects keep their positions and size.\nA smaller canvas may place objects outside the slide; Undo restores it.'))
        buttons=QDialogButtonBox(QDialogButtonBox.Ok|QDialogButtonBox.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);layout.addRow(buttons)
        if dialog.exec()==QDialog.Accepted:self.set_slide_size(width.value(),height.value())
    def play_video(self,shape):
        blob=movie_blob(shape)
        if blob is None:return
        from watchit import WatchIT
        if not hasattr(self,'media_temp'):self.media_temp=tempfile.TemporaryDirectory(prefix='freeconvert-slides-')
        path=Path(self.media_temp.name)/f'{shape.shape_id}.mp4';path.write_bytes(blob)
        player=WatchIT();player.add_media([str(path)]);player.show();self.players.append(player)
    def open(self,path=None):
        if path is None:
            name,_=QFileDialog.getOpenFileName(self,'Open presentation','','Presentations (*.pptx *.dazzle)');path=Path(name) if name else None
        if not path or not self.discard_changes():return
        def read():
            path_obj=Path(path)
            if path_obj.suffix.lower()=='.pptx':deck=Presentation(path_obj)
            else:
                data=json.loads(path_obj.read_text(encoding='utf-8'));slides=data['slides'];styles=data.get('styles',[dict(SLIDE_COLORS) for _ in slides])
                if not isinstance(slides,list) or not slides or not all(isinstance(s,list) and len(s)==2 and all(isinstance(v,str) for v in s) for s in slides):raise ValueError('Invalid Dazzle slides')
                if not isinstance(styles,list) or len(styles)!=len(slides) or any(not isinstance(s,dict) or set(s)!=set(SLIDE_COLORS) or any(not isinstance(v,str) or not re.fullmatch('#[0-9a-fA-F]{6}',v) for v in s.values()) for s in styles):raise ValueError('Invalid slide colors')
                deck=Presentation(io.BytesIO(base64.b64decode(data['pptx'],validate=True))) if 'pptx' in data else self._from_legacy(slides,styles)
            if not len(deck.slides):raise ValueError('Presentation contains no slides.')
            self.deck=deck;self.index=0;self.history=[];self.future=[];self.imported_pptx=path_obj if path_obj.suffix.lower()=='.pptx' else None;self._loaded_path=self.imported_pptx;self.path=None if self.imported_pptx else path_obj;self._native_changed=False;self.refresh();self.original_slides=list(self.slides);self.dirty=False
        self.report(read)
    def _write(self,path):
        path=Path(path)
        if self.imported_pptx and self.imported_pptx!=self._loaded_path:
            # Compatibility for callers supplying an original deck and edited title/body.
            pending=(self.title.toPlainText(),self.body.toPlainText());self.deck=Presentation(self.imported_pptx);self._loaded_path=self.imported_pptx;self.index=0;self._native_changed=False;self.refresh()
            with QSignalBlocker(self.title),QSignalBlocker(self.body):self.title.setPlainText(pending[0]);self.body.setPlainText(pending[1])
        self._store()
        if self.imported_pptx and path.resolve()==Path(self.imported_pptx).resolve() and self.path is None:raise ValueError('Save an imported presentation to a new file')
        if path.suffix.lower() not in ('.pptx','.dazzle'):raise ValueError('Choose PPTX or Dazzle format')
        if path.suffix.lower()=='.pptx' and self.imported_pptx and not self._native_changed:
            if path.resolve()!=Path(self.imported_pptx).resolve():shutil.copyfile(self.imported_pptx,path)
        else:
            blob=self._bytes()
            if path.suffix.lower()=='.dazzle':blob=json.dumps({'slides':self.slides,'styles':self.slide_styles,'pptx':base64.b64encode(blob).decode()},ensure_ascii=False).encode()
            with tempfile.NamedTemporaryFile(dir=path.parent,suffix=path.suffix,delete=False) as stream:stream.write(blob);temporary=Path(stream.name)
            try:temporary.replace(path)
            finally:temporary.unlink(missing_ok=True)
        self.path=path;self._native_changed=False;self.dirty=False
        if path.suffix.lower()=='.pptx':self.imported_pptx=self._loaded_path=path;self.original_slides=list(self.slides)
    def export_pdf(self):
        # Use the same canvas for print and editing, including inserted objects.
        import office_suite
        path=office_suite.pick_save(self,'Export slides','PDF (*.pdf)')
        if not path:return
        def write():
            from PySide6.QtGui import QPdfWriter, QPageSize, QPageLayout
            from PySide6.QtCore import QSizeF, QMarginsF
            writer=QPdfWriter(str(path));writer.setPageSize(QPageSize(QSizeF(self.deck.slide_width/EMU,self.deck.slide_height/EMU),QPageSize.Point));writer.setPageMargins(QMarginsF(0,0,0,0),QPageLayout.Point);writer.setResolution(72)
            self._store();painter=QPainter(writer)
            try:
                for i in range(len(self.deck.slides)):
                    if i:writer.newPage()
                    scene=self.make_scene(i,False);scene.render(painter,QRectF(0,0,writer.width(),writer.height()),scene.sceneRect());scene.deleteLater()
            finally:painter.end()
        self.report(write)
